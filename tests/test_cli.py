"""CLI-level regression tests — behavior that only exists at the `main()`
call site (flag wiring, unconditional post-extraction steps), not inside any
single extractor or relations function."""
from __future__ import annotations

import json
from pathlib import Path

from articling.cli import main


def test_numbered_headings_are_nested_without_vlm_enrichment(tmp_path: Path) -> None:
    """`nest_numbered_headings` needs no VLM/API key (see its own docstring),
    so the CLI must apply it unconditionally right after extraction — a user
    who never passes `--vlm-enrichment` should still get "3.1"/"3.2.1" nested
    under "3"/"3.2" whenever the source itself gives no native outline level
    for it (here: none of the paragraphs use Word's "Heading N" styles at
    all — someone typed the section numbers as plain body text instead — so
    `docx.py`'s own outline-based nesting has nothing to go on and every
    paragraph lands as a flat Artifact child; only the numbered-heading text
    pattern can recover the hierarchy). Regression for the 2026-09-13 change
    that promoted this out of the `--vlm-enrichment`-only bundle."""
    from docx import Document

    source = Document()
    source.add_paragraph("3 Results")
    source.add_paragraph("3.1 Analysis")
    source.add_paragraph("3.2 Discussion")
    source.add_paragraph("3.2.1 Detail")
    docx_path = tmp_path / "numbered.docx"
    source.save(docx_path)

    output_path = tmp_path / "out.json"
    assert main([str(docx_path), "-o", str(output_path)]) == 0

    data = json.loads(output_path.read_text(encoding="utf-8"))

    def node_for(text: str) -> dict:
        return next(n for n in data["nodes"] if n["properties"].get("text") == text)

    def parent_of(node_id: str) -> dict:
        return next(e for e in data["edges"] if e["type"] == "PARENT_OF" and e["target_id"] == node_id)

    section = node_for("3 Results")
    subsection = node_for("3.2 Discussion")

    for subheading_text, expected_parent in (
        ("3.1 Analysis", section),
        ("3.2 Discussion", section),
        ("3.2.1 Detail", subsection),
    ):
        subheading = node_for(subheading_text)
        parent_edge = parent_of(subheading["id"])
        assert parent_edge["source_id"] == expected_parent["id"]
        assert parent_edge["properties"].get("nested_by") == "numbered_heading_pattern"
