"""Generates synthetic DOCX/PPTX/XLSX files for testing, in code.

Built fresh each time so no real company document ever gets committed to
the repo (an open-source premise) — each file always includes at least one
Text + Table + Image, so all three extractors can be verified to produce
all three Node types.

Some fixtures below deliberately use Korean text (e.g. "표 1." / "그림 1."
caption prefixes, Korean cell values) — that's real test data exercising
the deterministic Korean caption-prefix heuristics
(`_CAPTION_PREFIXES`) and Korean-text handling, not something to translate.
"""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import docx
import openpyxl
import pymupdf
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Border, Side
from PIL import Image as PILImage
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, PROG_ID
from pptx.util import Inches


def _tiny_png_bytes(color: tuple[int, int, int] = (200, 60, 60)) -> bytes:
    buf = io.BytesIO()
    PILImage.new("RGB", (40, 30), color).save(buf, format="PNG")
    return buf.getvalue()


def build_docx(path: Path) -> Path:
    d = docx.Document()
    d.add_paragraph("샘플 보고서 제목")
    d.add_paragraph("표 1. 측정 결과")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "값"
    table.cell(1, 0).text = "토크"
    table.cell(1, 1).text = "3.4 kgf"
    d.add_paragraph("그림 1. 시료 사진")
    img_path = path.parent / "_tmp_docx_img.png"
    img_path.write_bytes(_tiny_png_bytes())
    d.add_picture(str(img_path))
    d.save(str(path))
    img_path.unlink(missing_ok=True)
    return path


def build_docx_with_table_cell_image(path: Path) -> Path:
    """A DOCX with an image inside a table cell, followed by more
    paragraphs — for a regression test verifying that `extractors/docx.py`
    absorbs a table-cell image into that Table's `embedded_images`
    properties (the same "Option A: absorb" rule as xlsx.py) instead of
    creating a separate Image node for it."""
    d = docx.Document()
    d.add_paragraph("표 1. 셀 안에 사진이 있는 표")
    table = d.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    img_path = path.parent / "_tmp_docx_cell_img.png"
    img_path.write_bytes(_tiny_png_bytes((90, 60, 200)))
    cell.paragraphs[0].add_run().add_picture(str(img_path))
    d.add_paragraph("표 바로 다음 문단")
    d.add_paragraph("그 다음 문단")
    d.add_paragraph("문서 맨 끝 문단")
    d.save(str(path))
    img_path.unlink(missing_ok=True)
    return path


def build_pptx(path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(4), Inches(0.5))
    tb.text_frame.text = "표 1. 요약"

    rows, cols = 2, 2
    graphic_frame = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(1), Inches(3), Inches(1))
    table = graphic_frame.table
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "값"
    table.cell(1, 0).text = "토크"
    table.cell(1, 1).text = "3.4 kgf"

    img_path = path.parent / "_tmp_pptx_img.png"
    img_path.write_bytes(_tiny_png_bytes((60, 120, 200)))
    slide.shapes.add_picture(str(img_path), Inches(0.5), Inches(2.5), height=Inches(1))
    img_path.unlink(missing_ok=True)

    prs.save(str(path))
    return path


def build_pptx_with_ole(path: Path) -> Path:
    """A slide holding one embedded OLE object (e.g. an Excel worksheet
    dropped onto a slide, shown as an icon) — reproduces the real-document
    case `pptx.py`'s `_OLE_SHAPE_TYPES` handling exists for: this shape type
    has no python-pptx picture API at all (`shape.image` doesn't exist on
    it), so without special-casing it, `_classify_shape` falls through every
    branch and the shape silently disappears from the graph."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    xlsx_bytes = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.active["A1"] = "임베디드 시트"
    wb.save(xlsx_bytes)
    xlsx_bytes.seek(0)

    icon_bytes = io.BytesIO(_tiny_png_bytes((200, 200, 60)))
    slide.shapes.add_ole_object(
        xlsx_bytes, PROG_ID.XLSX, Inches(1), Inches(1),
        icon_file=icon_bytes, icon_width=Inches(1), icon_height=Inches(1),
    )

    prs.save(str(path))
    return path


def build_pptx_with_box_over_image(path: Path) -> Path:
    """A picture with a plain, textless rectangle drawn directly on top of
    it (PowerPoint's "insert shape > rectangle", no fill text) — reproduces
    the real-document case `pptx.py`'s `_render_slide_crops` exists for:
    the rectangle has no text of its own, so `_classify_shape`
    never turns it into a Node, and the picture's saved bytes are its
    original embedded file (not a slide render), so without that fix the
    box disappears from the graph entirely."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(8), Inches(6)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    img_path = path.parent / "_tmp_pptx_box_img.png"
    img_path.write_bytes(_tiny_png_bytes((0, 120, 200)))
    slide.shapes.add_picture(str(img_path), Inches(1), Inches(1), Inches(4), Inches(3))
    img_path.unlink(missing_ok=True)

    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(1.5), Inches(1), Inches(1))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 0, 0)
    box.line.fill.background()

    prs.save(str(path))
    return path


def build_pptx_with_shaded_table(path: Path) -> Path:
    """A native table with a colored header row — reproduces the real-
    document case `pptx.py`'s `_render_slide_crops` table-capture handling
    exists for: cell shading has nowhere to live in the plain-text `grid`
    (`cell.text` only), so without a rendered `capture_path`, PowerPoint's
    visual formatting (which the reader relies on to tell a header row from
    data) is silently lost — unlike XLSX tables, which already get one."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(8), Inches(6)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    graphic_frame = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    table = graphic_frame.table
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "값"
    table.cell(1, 0).text = "토크"
    table.cell(1, 1).text = "3.4 kgf"
    for col in range(2):
        cell = table.cell(0, col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(255, 0, 0)

    prs.save(str(path))
    return path


def build_pptx_with_chart(path: Path) -> Path:
    """A slide holding one native chart (a `p:graphicFrame` with chart
    graphicData) — reproduces the real-document case `pptx.py`'s chart
    handling exists for: `GraphicFrame.shape_type` reports
    `MSO_SHAPE_TYPE.CHART` for this, which `_classify_shape` didn't check
    for at all before, so the chart (and its data) silently disappeared."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    chart_data = CategoryChartData()
    chart_data.categories = ["1분기", "2분기", "3분기"]
    chart_data.add_series("매출", (19.2, 21.4, 16.7))
    chart_data.add_series("비용", (10.1, 12.3, 9.9))
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(4), chart_data,
    )
    graphic_frame.chart.has_title = True
    graphic_frame.chart.chart_title.text_frame.text = "분기별 실적"

    prs.save(str(path))
    return path


def build_pptx_with_smartart(path: Path) -> Path:
    """A slide holding one SmartArt diagram (a `p:graphicFrame` whose
    graphicData URI is the diagram namespace) — reproduces the real-
    document case `pptx.py`'s SmartArt handling exists for.

    python-pptx has no authoring API for SmartArt at all, so unlike the
    other `build_pptx_with_*` fixtures this can't be built by calling
    shape-adding methods — the diagram data part, its content-type
    registration, its relationship from the slide, and the `p:graphicFrame`
    referencing it via `<dgm:relIds>` all have to be injected as raw XML
    into a saved package, the same technique `_inject_into_drawing_xml`
    uses for XLSX shapes python-pptx/openpyxl can't author either. One
    `type="doc"` point (the diagram's root, never user-facing) and one
    `type="parTrans"` point (a transition/connector point with no real
    text) are included specifically to exercise `_smartart_text`'s
    filtering — only the two `type="node"` points' text should come out."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    base_path = path.parent / "_tmp_pptx_smartart_base.pptx"
    prs.save(str(base_path))

    with zipfile.ZipFile(base_path) as zin:
        entries = {name: zin.read(name) for name in zin.namelist()}
    base_path.unlink(missing_ok=True)

    diagram_ns = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    entries["ppt/diagrams/data1.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<dgm:dataModel xmlns:dgm="{diagram_ns}" xmlns:a="{a_ns}"><dgm:ptLst>'
        f'<dgm:pt modelId="0" type="doc"><dgm:t><a:bodyPr/><a:lstStyle/>'
        f'<a:p><a:r><a:t>ROOT_SHOULD_BE_SKIPPED</a:t></a:r></a:p></dgm:t></dgm:pt>'
        f'<dgm:pt modelId="1" type="node"><dgm:t><a:bodyPr/><a:lstStyle/>'
        f'<a:p><a:r><a:t>1단계: 준비</a:t></a:r></a:p></dgm:t></dgm:pt>'
        f'<dgm:pt modelId="2" type="node"><dgm:t><a:bodyPr/><a:lstStyle/>'
        f'<a:p><a:r><a:t>2단계: 조립</a:t></a:r></a:p></dgm:t></dgm:pt>'
        f'<dgm:pt modelId="3" type="parTrans"><dgm:t><a:bodyPr/><a:lstStyle/>'
        f'<a:p><a:endParaRPr/></a:p></dgm:t></dgm:pt>'
        f'</dgm:ptLst></dgm:dataModel>'
    ).encode()

    content_types = entries["[Content_Types].xml"].decode()
    override = (
        '<Override PartName="/ppt/diagrams/data1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml"/>'
    )
    entries["[Content_Types].xml"] = content_types.replace("</Types>", override + "</Types>").encode()

    rels_name = "ppt/slides/_rels/slide1.xml.rels"
    rels = entries[rels_name].decode()
    new_rel = (
        '<Relationship Id="rIdDgm" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData" '
        'Target="../diagrams/data1.xml"/>'
    )
    entries[rels_name] = rels.replace("</Relationships>", new_rel + "</Relationships>").encode()

    slide_xml = entries["ppt/slides/slide1.xml"].decode()
    graphic_frame = (
        '<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="99" name="SmartArt 1"/>'
        '<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        '<p:xfrm><a:off x="1000000" y="1000000"/><a:ext cx="3000000" cy="2000000"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="{diagram_ns}">'
        f'<dgm:relIds xmlns:dgm="{diagram_ns}" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'r:dm="rIdDgm" r:lo="rIdDgm" r:qs="rIdDgm" r:cs="rIdDgm"/>'
        '</a:graphicData></a:graphic></p:graphicFrame>'
    )
    entries["ppt/slides/slide1.xml"] = slide_xml.replace("</p:spTree>", graphic_frame + "</p:spTree>").encode()

    with zipfile.ZipFile(path, "w") as zout:
        for name, data in entries.items():
            zout.writestr(name, data)
    return path


def build_xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # a table closed by borders (tests the border-based detection path) — B2:C3
    thin = Side(style="thin")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    ws["B2"] = "항목"
    ws["C2"] = "값"
    ws["B3"] = "토크"
    ws["C3"] = "3.4 kgf"
    for row in ws["B2:C3"]:
        for cell in row:
            cell.border = border

    # an image inside the table range (B2:C3) — an "Option A" absorption target
    img_bytes = _tiny_png_bytes((60, 200, 90))
    img_path = path.parent / "_tmp_xlsx_img_in_table.png"
    img_path.write_bytes(img_bytes)
    in_table_img = XLImage(str(img_path))
    in_table_img.anchor = "B2"
    ws.add_image(in_table_img)

    # a standalone image outside the table range — must become a separate Image node
    img_path2 = path.parent / "_tmp_xlsx_img_standalone.png"
    img_path2.write_bytes(_tiny_png_bytes((200, 160, 40)))
    standalone_img = XLImage(str(img_path2))
    standalone_img.anchor = "F10"
    ws.add_image(standalone_img)

    # a note physically separate from the table — must become its own Text connected-component
    ws["B10"] = "메모: 재측정 필요"

    wb.save(str(path))
    img_path.unlink(missing_ok=True)
    img_path2.unlink(missing_ok=True)
    return path


def build_xlsx_with_overlapping_standalone_images(path: Path) -> Path:
    """Two standalone images overlapping each other, outside any table
    range — reproduces a real case (a driver-icon photo placed separately
    on top of a frame photo in a screw fastening-strength review document)
    of a small icon photo placed on top of a large background photo. Both
    are anchored at F10 (the same starting pixel) to guarantee overlap —
    `extract()` should group these two by overlap, keeping only the larger
    one (the background photo) as an Image node and compositing the
    smaller one (the icon) onto it, absorbing it (the same "Option A:
    absorb" principle, just with another standalone image as the
    absorption target instead of a table)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    base_buf = io.BytesIO()
    PILImage.new("RGB", (200, 150), (200, 160, 40)).save(base_buf, format="PNG")
    base_path = path.parent / "_tmp_xlsx_img_overlap_base.png"
    base_path.write_bytes(base_buf.getvalue())
    base_img = XLImage(str(base_path))
    base_img.anchor = "F10"
    ws.add_image(base_img)

    overlay_path = path.parent / "_tmp_xlsx_img_overlap_overlay.png"
    overlay_path.write_bytes(_tiny_png_bytes((30, 30, 200)))
    overlay_img = XLImage(str(overlay_path))
    overlay_img.anchor = "F10"  # the same anchor -> starts at the same pixel as the background photo, guaranteed overlap
    ws.add_image(overlay_img)

    wb.save(str(path))
    base_path.unlink(missing_ok=True)
    overlay_path.unlink(missing_ok=True)
    return path


def build_xlsx_with_table_caption(path: Path) -> Path:
    """An XLSX with caption text directly above a table — for a regression
    test of `relations/propose.py`'s XLSX layout crop (a crop showing how
    the table+caption are actually laid out on the sheet)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws["B2"] = "표 1. 측정 결과"

    thin = Side(style="thin")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    ws["B4"] = "항목"
    ws["C4"] = "값"
    ws["B5"] = "토크"
    ws["C5"] = "3.4 kgf"
    for row in ws["B4:C5"]:
        for cell in row:
            cell.border = border

    wb.save(str(path))
    return path


def build_pdf(path: Path, n_pages: int = 1) -> Path:
    """A synthetic PDF with text + image blocks. If `n_pages>1`, repeats a
    heading+body on each page — used for testing text/image extraction
    across multiple pages in extractors/pdf.py (reading order, etc.).

    Uses English text only — pymupdf's default built-in font (Helvetica)
    has no Korean glyphs, so putting Korean into `insert_text` breaks into
    replacement characters (`··`). Real PDFs mostly embed Korean glyphs, so
    this is a constraint specific to this test fixture only."""
    doc = pymupdf.open()
    for i in range(n_pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"Section {i + 1}. Heading", fontsize=16)
        page.insert_text((72, 100), f"Page {i + 1} body content goes here.", fontsize=11)
        if i == 0:
            page.insert_image(
                pymupdf.Rect(72, 140, 172, 200), stream=_tiny_png_bytes((90, 60, 200))
            )
    doc.save(str(path))
    doc.close()
    return path


def build_pdf_out_of_stream_order(path: Path) -> Path:
    """A one-page PDF where the order written to the content stream is
    deliberately reversed from the on-screen (top-to-bottom) position —
    "Bottom" is inserted first, but given coordinates that put it below
    "Top" on screen. For a regression test verifying that
    `extractors/pdf.py` restores the actual visual order with
    `page.get_text("dict", sort=True)` (see the empirical confirmation in
    `docs/docling-compat-harness.md` §4a — a document exported from PPT
    really did have its stream order reversed like this)."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 400), "Bottom paragraph (inserted first)", fontsize=11)
    page.insert_text((72, 72), "Top paragraph (inserted second)", fontsize=11)
    doc.save(str(path))
    doc.close()
    return path


def build_pdf_with_vector_table(path: Path) -> Path:
    """A one-page PDF with a 2x2 table actually drawn with vector lines (3
    horizontal + 3 vertical) — for a regression test of
    `extractors/pdf_tables.detect_table_candidates`'s "boxed (grid)" mode
    and `relations/table_structure.enrich_pdf_tables`'s absorption/Table
    node creation. A paragraph is also placed outside the table so
    "content that isn't an absorption target is left untouched" can be
    verified at the same time."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 60), "Report heading (outside the table)", fontsize=12)

    x0, y0, cell_w, cell_h = 72, 100, 100, 40
    for row in range(3):
        page.draw_line((x0, y0 + row * cell_h), (x0 + 2 * cell_w, y0 + row * cell_h))
    for col in range(3):
        page.draw_line((x0 + col * cell_w, y0), (x0 + col * cell_w, y0 + 2 * cell_h))

    cells = [["Item", "Value"], ["Torque", "3.4 kgf"]]
    for r in range(2):
        for c in range(2):
            page.insert_text((x0 + c * cell_w + 10, y0 + r * cell_h + 25), cells[r][c], fontsize=10)

    doc.save(str(path))
    doc.close()
    return path


def build_pdf_with_vector_figure(path: Path) -> Path:
    """A one-page PDF that draws one vector figure out of hundreds of small
    filled rectangles (attention-heatmap style) with a "Figure 1. …"
    caption right below it — for a regression test of
    `extractors/pdf_figures.detect_vector_figure_regions`/`enrich_pdf_figures`
    (empirical basis: `1706.03762` Figure 3/4/5, see the module docstring).
    Also places a decorative line unrelated to the figure (a single
    rectangle border, only 4 items) so "low-density vector graphics aren't
    mistaken for a figure" can be verified at the same time."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 60), "Report heading (outside the figure)", fontsize=12)

    # a 12x10 grid of small filled rectangles = 120 items (exceeds _MIN_ITEMS_FOR_FIGURE=100)
    cell = 10
    x0, y0 = 72, 100
    for row in range(10):
        for col in range(12):
            rect = pymupdf.Rect(x0 + col * cell, y0 + row * cell, x0 + (col + 1) * cell, y0 + (row + 1) * cell)
            page.draw_rect(rect, color=(0.6, 0.6, 0.6), fill=(0.6, 0.2, 0.8), width=0.5)

    page.insert_text((72, y0 + 10 * cell + 20), "Figure 1. Synthetic attention heatmap", fontsize=11)

    # a single decorative border rectangle unrelated to the figure (only 4 segments) — must not become a figure candidate due to low density
    page.draw_rect(pymupdf.Rect(72, 500, 300, 540), color=(0, 0, 0), width=1)

    doc.save(str(path))
    doc.close()
    return path


def build_pdf_with_figure_caption(path: Path) -> Path:
    """A one-page PDF where a "Figure 1. …" caption comes right after an
    image — for a regression test of extractors/pdf.py's deterministic
    CAPTION_OF heuristic (`_CAPTION_PREFIXES`)."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(72, 72, 172, 132), stream=_tiny_png_bytes((90, 60, 200)))
    page.insert_text((72, 140), "Figure 1. Sample photo", fontsize=11)
    doc.save(str(path))
    doc.close()
    return path
