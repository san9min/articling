"""Regressions from grouped supplier slides: pixels, paint order and text fit."""
from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from articling.capture.pptx_capture import PptxSlideRenderer
from articling.extractors.pptx import extract
from articling.scaffold import check_invariants


def _deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(8), Inches(6)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _rect(shapes, x, y, w, h, color):
    shape = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*color)
    shape.line.fill.background()
    return shape


def test_native_crop_alpha_paint_order_and_group_transform(tmp_path):
    prs, slide = _deck()
    _rect(slide.shapes, 0, 0, 8, 6, (255, 255, 0))
    # Transparent image has red left half, blue top-right, transparent bottom-right.
    image = Image.new('RGBA', (100, 100), (255, 0, 0, 255))
    image.paste((0, 0, 255, 255), (50, 0, 100, 50))
    image.paste((0, 0, 0, 0), (50, 50, 100, 100))
    source = tmp_path / 'rgba.png'
    image.save(source)
    group = slide.shapes.add_group_shape()
    picture = group.shapes.add_picture(str(source), Inches(1), Inches(1), Inches(2), Inches(2))
    picture.crop_left = 0.5
    group.left, group.top, group.width, group.height = Inches(4), Inches(1), Inches(1), Inches(2)
    # Added later but positioned above the picture: reading-order painting
    # would incorrectly paint this green foreground underneath the picture.
    _rect(slide.shapes, 4.5, 0.5, 1, 1, (0, 255, 0))
    path = tmp_path / 'source.pptx'
    prs.save(path)
    renderer = PptxSlideRenderer(path, max_dimension=800)
    result = renderer.render(0)
    pixels = Image.open(BytesIO(result.png))
    assert pixels.size == (800, 600)
    assert pixels.getpixel((425, 175)) == (0, 0, 255)  # Crop removed the red half.
    assert pixels.getpixel((425, 250)) == (255, 255, 0)  # Alpha preserved background.
    assert pixels.getpixel((500, 125)) == (0, 255, 0)  # XML paint order, not reading order.
    assert pixels.getpixel((150, 150)) == (255, 255, 0)  # No image at stale child coordinates.
    doc = extract(path, capture_dir=tmp_path / 'captures')
    assert check_invariants(doc.nodes, doc.edges) == []


def test_picture_rotation_and_reflection_affect_pixels(tmp_path):
    prs, slide = _deck()
    source = tmp_path / 'two-colors.png'
    image = Image.new('RGB', (200, 100), 'red')
    image.paste((0, 0, 255), (100, 0, 200, 100))
    image.save(source)
    picture = slide.shapes.add_picture(str(source), Inches(2), Inches(2), Inches(2), Inches(1))
    picture.rotation = 90
    picture._element.spPr.xfrm.set('flipH', '1')
    path = tmp_path / 'rotated.pptx'
    prs.save(path)
    renderer = PptxSlideRenderer(path, max_dimension=800)
    pixels = Image.open(BytesIO(renderer.render(0).png))
    assert pixels.getpixel((300, 200)) == (0, 0, 255)
    assert pixels.getpixel((300, 300)) == (255, 0, 0)
    assert pixels.getpixel((225, 250)) == (255, 255, 255)


def test_merged_table_and_multiline_text_stay_inside_cells(tmp_path):
    prs, slide = _deck()
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
    table.rows[0].height, table.rows[1].height = Inches(0.5), Inches(1.5)
    table.cell(0, 0).merge(table.cell(0, 1))
    title = table.cell(0, 0)
    title.fill.solid()
    title.fill.fore_color.rgb = RGBColor(255, 255, 0)
    title.text = 'FIRST\nSECOND'
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(24)
            run.font.color.rgb = RGBColor(255, 0, 0)
    for ci in range(2):
        cell = table.cell(1, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0, 0, 255)
    path = tmp_path / 'table.pptx'
    prs.save(path)
    result = PptxSlideRenderer(path, max_dimension=800).render(0)
    pixels = Image.open(BytesIO(result.png))
    # The merge's right half retains its fill. Both fitted text lines survive;
    # red glyph pixels never spill into the next blue row.
    assert pixels.getpixel((400, 125)) == (255, 255, 0)
    red_ys = {y for y in range(100, 150) for x in range(100, 300)
              if (lambda c: c[0] > 180 and c[1] < 100 and c[2] < 100)(pixels.getpixel((x, y)))}
    assert len(red_ys) >= 10
    assert any(y < 125 for y in red_ys) and any(y >= 125 for y in red_ys)
    assert pixels.getpixel((120, 170)) == (0, 0, 255)
    assert any('text fitted' in warning for warning in result.warnings)


def test_bad_picture_skips_only_that_shape(tmp_path):
    prs, slide = _deck()
    source = tmp_path / 'image.png'
    Image.new('RGB', (20, 20), 'red').save(source)
    picture = slide.shapes.add_picture(str(source), Inches(1), Inches(1))
    picture.part.related_part(picture._element.blip_rId)._blob = b'invalid bitmap'
    _rect(slide.shapes, 4, 1, 1, 1, (0, 255, 0))
    path = tmp_path / 'bad-picture.pptx'
    prs.save(path)
    result = PptxSlideRenderer(path, max_dimension=800).render(0)
    assert result.warnings and any('skipped' in warning for warning in result.warnings)
    assert Image.open(BytesIO(result.png)).getpixel((450, 150)) == (0, 255, 0)
