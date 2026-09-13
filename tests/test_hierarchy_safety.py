"""Numbering is meaningful only inside its native structural scope."""
from pathlib import Path

from articling.relations.propose import _apply_heading_reparents, _enumerated_sibling_groups, nest_numbered_headings
from articling.schema import ArticDocument, Edge, EdgeType, Node, NodeType
from articling.scaffold import check_invariants


def text(identity: str, value: str, **properties) -> Node:
    return Node(id=identity, type=NodeType.TEXT, name=value, properties={"text": value, **properties})


def document(*sections: list[Node]) -> ArticDocument:
    nodes, edges = [], []
    for i, section in enumerate(sections):
        artifact = Node(id=f"a{i}", type=NodeType.ARTIFACT, name=f"Section {i}")
        nodes.extend([artifact, *section])
        edges.extend(Edge(type=EdgeType.PARENT_OF, source_id=artifact.id, target_id=n.id) for n in section)
    return ArticDocument(source_path="test", format="pptx", nodes=nodes, edges=edges)


def test_numbering_cannot_borrow_parent_from_another_slide_or_closed_section():
    doc = document([text("first", "1 Overview")], [text("orphan", "1.1 Details")], [
        text("old", "1 Earlier"), text("oldsub", "1.1 Earlier detail"),
        text("new", "1 Restart"), text("newsub", "1.1 Current detail"),
        text("second", "2 Next"), text("late", "1.2 Unmatched"),
    ])
    assert nest_numbered_headings(doc) == ["oldsub", "newsub"]
    parents = {e.target_id: e.source_id for e in doc.edges}
    assert parents["orphan"] == "a1"
    assert parents["newsub"] == "new"
    assert parents["late"] == "a2"
    assert check_invariants(doc.nodes, doc.edges) == []


def test_heading_mutations_preserve_artifacts_and_prevent_self_and_ancestor_cycles():
    parent, child = text("p", "1 Parent"), text("c", "1.1 Child")
    remote = text("r", "Remote")
    doc = document([parent, child], [remote])
    # Existing model mistake has reversed the expected numeric hierarchy.
    doc.edges[0] = Edge(type=EdgeType.PARENT_OF, source_id=child.id, target_id=parent.id)
    original = list(doc.edges)
    assert nest_numbered_headings(doc) == []
    assert _apply_heading_reparents(doc, "fake", [(child, remote, "cross slide"), (child, child, "self"), (child, parent, "cycle")]) == []
    assert doc.edges == original
    assert check_invariants(doc.nodes, doc.edges) == []
    doc.edges.append(Edge(type=EdgeType.PARENT_OF, source_id=parent.id, target_id=child.id))
    problems = check_invariants(doc.nodes, doc.edges)
    assert any("PARENT_OF cycle" in problem for problem in problems)
    assert any("fan-in" in problem for problem in problems)


def test_native_lists_do_not_bridge_heading_or_parent_boundaries():
    items = [text(f"i{i}", f"Item {i}", list_num_id=7, list_level=0) for i in range(6)]
    heading = text("h", "New section", outline_level=0)
    doc = document([*items[:2], heading, *items[2:]])
    doc.format = "docx"
    for edge in doc.edges:
        if edge.target_id in {"i4", "i5"}:
            edge.source_id = heading.id
    assert [[n.id for n in group] for group in _enumerated_sibling_groups(doc)] == [["i0", "i1"], ["i2", "i3"], ["i4", "i5"]]


def test_docx_zero_num_id_removes_numbering(tmp_path: Path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from articling.extractors.docx import extract

    source = Document()
    for identity in (7, 0, 0):
        paragraph = source.add_paragraph(f"Paragraph {identity}")
        numbering = OxmlElement("w:numPr")
        for tag, value in (("w:ilvl", "0"), ("w:numId", str(identity))):
            element = OxmlElement(tag)
            element.set(qn("w:val"), value)
            numbering.append(element)
        paragraph._p.get_or_add_pPr().append(numbering)
    path = tmp_path / "numbering-removal.docx"
    source.save(path)
    doc = extract(path, capture_dir=tmp_path / "captures")
    paragraphs = [n for n in doc.nodes if n.type == NodeType.TEXT]
    assert paragraphs[0].properties["list_num_id"] == 7
    assert all("list_num_id" not in n.properties for n in paragraphs[1:])
    assert _enumerated_sibling_groups(doc) == []
    assert check_invariants(doc.nodes, doc.edges) == []
