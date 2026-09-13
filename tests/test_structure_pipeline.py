"""Native placement survives semantic hierarchy and physical page changes."""
from pathlib import Path

import pymupdf
from docx import Document

from articling.extractors import docx, pdf
from articling.schema import EdgeType, NodeType
from articling.scaffold import check_invariants
from articling.structure import finalize_structure, content_artifact_map


def test_docx_page_break_does_not_close_section(tmp_path: Path):
    source = Document()
    source.add_heading("Section", 1)
    source.add_page_break()
    source.add_paragraph("Continued body")
    source.add_heading("Next section", 1)
    source.add_paragraph("Next body")
    path = tmp_path / "pages.docx"
    source.save(path)
    doc = docx.extract(path, capture_dir=tmp_path / "captures")
    texts = {n.properties.get("text"): n for n in doc.content_nodes()}
    parents = {e.target_id: e.source_id for e in doc.edges if e.type == EdgeType.PARENT_OF}
    assert parents[texts["Continued body"].id] == texts["Section"].id
    assert parents[texts["Next body"].id] == texts["Next section"].id
    artifact = next(n for n in doc.nodes if n.type == NodeType.ARTIFACT)
    assert texts["Continued body"].properties["native_location"] == {"artifact_id": artifact.id}
    assert check_invariants(doc.nodes, doc.edges) == []


def test_pdf_cross_page_outline_keeps_missing_ancestor_barrier(tmp_path: Path):
    path = tmp_path / "pages.pdf"
    with pymupdf.open() as source:
        for title in ("Root", "Child", "Orphan"):
            source.new_page().insert_text((72, 72), title)
        source.set_toc([[1, "Root", 1], [2, "Child", 2],
                        [1, "Missing", 3], [2, "Orphan", 3]])
        source.save(path)
    doc = pdf.extract(path, capture_dir=tmp_path / "captures")
    texts = {n.properties.get("text"): n for n in doc.content_nodes()}
    artifact = next(n for n in doc.nodes if n.type == NodeType.ARTIFACT)
    parents = {e.target_id: e.source_id for e in doc.edges if e.type == EdgeType.PARENT_OF}
    assert parents[texts["Child"].id] == texts["Root"].id
    assert parents[texts["Orphan"].id] == artifact.id
    assert texts["Child"].properties["native_location"] == {"artifact_id": artifact.id, "page_index": 1}
    before = doc.model_dump()
    finalize_structure(doc)
    assert doc.model_dump() == before
    # Native scope remains available even if a semantic edge is removed.
    doc.edges = [e for e in doc.edges if e.target_id != texts["Child"].id]
    assert content_artifact_map(doc)[texts["Child"].id] == artifact.id
    assert check_invariants(doc.nodes, doc.edges) == []
