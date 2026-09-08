"""PPTX -> ArticDocument, a fully native path (python-pptx only, no LibreOffice/VLM needed).

One slide = one Artifact. Shapes within each slide are sorted by visual bbox (top, left)
to approximate reading order — python-pptx's shape order is z-order/insertion
order, which can differ from actual reading order. Group shapes are flattened
recursively, with ancestor transforms retained in the slide-space bbox.

Node: Text (text frame), Table (cell grid — also a chart's categories/series
      data, see below), Image (picture shape). Every content node preserves
      its slide_index, original top/left/width/height in its parent
      coordinate system, and a 0-1000 normalized slide-space bbox.
Edge: PARENT_OF, NEXT (order between slides only — never links content
      within one slide), CAPTION_OF (the same heuristic as docx.py),
      REFERENCES (a connector/arrow shape — never a Node itself, no schema
      type fits a bare line — glued to two other shapes via PowerPoint's
      `stCxn`/`endCxn`, or, absent that, geometrically touching them; see
      `_connector_reference_edge`)

A shape with no text of its own that isn't a table/chart/OLE object/picture
(a highlight box, an unglued arrow) never becomes a Node either — but if it
overlaps a picture, it would otherwise vanish from that picture's saved
bytes too, since those are the original embedded file, not a slide render.
`_render_slide_crops` composites it onto that picture instead (re-rendering
just that picture's region via `capture.pptx_capture.PptxSlideRenderer`),
the same fix `capture.xlsx_capture.annotate_standalone_image` applies to
XLSX's highlight-box-on-photo case. The same function also gives every
native table a `capture_path` — a visual capture PPTX tables otherwise
never had (unlike XLSX's), since `PptxSlideRenderer` already draws cell
shading/merges/borders for slide reconstruction anyway.

A chart (`GraphicFrame.shape_type == MSO_SHAPE_TYPE.CHART`) becomes a Table
node — its categories/series data, not a re-rendered picture (this repo has
no chart-drawing code, and the exact data is more useful than an approximate
picture anyway); see `_chart_grid`. SmartArt has no python-pptx object model
and no rendering support at all (`shape_type` reports `None` for it, since
it's neither table/chart/OLE) — every text label typed into it is still
recovered from the diagram's data-model part and becomes a Text node; see
`_smartart_text`.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import math

from lxml import etree
from PIL import Image as PILImage
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from ..schema import ArticDocument, Edge, EdgeType, Node, NodeType
from ..scaffold import caption_prefix_edges, file_node, parent_edges, resolve_capture_dir, save_image_bytes

# An embedded/linked OLE object (e.g. an Excel worksheet dropped onto a
# slide) has no python-pptx picture API at all — without this, `_classify_shape`
# falls through every branch and the shape (and its position) disappears from
# the graph entirely (confirmed on a real document: two `Excel.Sheet` OLE
# objects on one slide, silently dropped). See `_ole_preview_image`.
_OLE_SHAPE_TYPES = (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT)
# SmartArt has no python-pptx object model at all: `GraphicFrame.shape_type`
# (graphfrm.py) returns `None` for it rather than raising or matching a known
# member, since it's neither table/chart/OLE — `_classify_shape` fell through
# every branch and dropped it entirely, same failure mode as the OLE case
# above. Detected by graphicData URI directly (`_graphic_data_uri`) rather
# than by `shape_type is None`, since that would also swallow any other
# not-yet-implemented graphic frame content under the same label.
_SMARTART_GRAPHIC_DATA_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_DIAGRAM_NS = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"


def _flatten_shapes(shapes) -> list:
    flat = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            flat.extend(_flatten_shapes(shape.shapes))
        else:
            flat.append(shape)
    return flat


def _shape_slide_corners(shape) -> list[tuple[float, float]]:
    """Map local TL/TR/BR/BL corners to slide EMU, retaining orientation.

    Shared by extraction and reconstruction so their geometry cannot diverge.
    """
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    # OOXML rotation is around the shape center. Keep raw EMU properties,
    # but use the rotated axis-aligned envelope for visual matching.
    rotation = shape.rotation
    angle = math.radians(rotation)
    cx, cy = left + width / 2, top + height / 2
    transform = shape._element.find(f"{qn('p:spPr')}/{qn('a:xfrm')}")
    flip_x = -1 if transform is not None and transform.get("flipH") in ("1", "true") else 1
    flip_y = -1 if transform is not None and transform.get("flipV") in ("1", "true") else 1
    points = []
    for dx, dy in [(-width / 2, -height / 2), (width / 2, -height / 2),
                   (width / 2, height / 2), (-width / 2, height / 2)]:
        dx, dy = dx * flip_x, dy * flip_y
        points.append((cx + dx * math.cos(angle) - dy * math.sin(angle),
                       cy + dx * math.sin(angle) + dy * math.cos(angle)))
    # Children use the group's chOff/chExt coordinate system, not slide EMU.
    # Transform actual corners through every ancestor before taking an envelope:
    # taking envelopes first inflates nested rotated groups.
    parent = shape._element.getparent()
    while parent is not None and parent.tag == qn("p:grpSp"):
        transform = parent.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
        if transform is not None:
            off, ext = transform.find(qn("a:off")), transform.find(qn("a:ext"))
            child_off, child_ext = transform.find(qn("a:chOff")), transform.find(qn("a:chExt"))
            if all(value is not None for value in (off, ext, child_off, child_ext)):
                ox, oy = int(off.get("x")), int(off.get("y"))
                ew, eh = int(ext.get("cx")), int(ext.get("cy"))
                cw, ch = int(child_ext.get("cx")), int(child_ext.get("cy"))
                sx, sy = ew / cw if cw else 1, eh / ch if ch else 1
                gx, gy = ox + ew / 2, oy + eh / 2
                radians = math.radians(int(transform.get("rot", "0")) / 60000)
                flip_x = -1 if transform.get("flipH") in ("1", "true") else 1
                flip_y = -1 if transform.get("flipV") in ("1", "true") else 1
                mapped = []
                for x, y in points:
                    dx = (ox + (x - int(child_off.get("x"))) * sx - gx) * flip_x
                    dy = (oy + (y - int(child_off.get("y"))) * sy - gy) * flip_y
                    mapped.append((gx + dx * math.cos(radians) - dy * math.sin(radians),
                                   gy + dx * math.sin(radians) + dy * math.cos(radians)))
                points = mapped
        parent = parent.getparent()
    return points


def _shape_layout_properties(shape, slide_idx: int, slide_width: int, slide_height: int) -> dict:
    """Keep raw parent-space EMU and the transformed slide-space envelope."""
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    rotation = shape.rotation
    points = _shape_slide_corners(shape)
    x0, y0 = min(x for x, y in points), min(y for x, y in points)
    x1, y1 = max(x for x, y in points), max(y for x, y in points)
    return {
        "slide_index": slide_idx,
        "top": top,
        "left": left,
        "width": width,
        "height": height,
        "rotation": rotation,
        "bbox": {
            "x_min": round(x0 / slide_width * 1000),
            "y_min": round(y0 / slide_height * 1000),
            "x_max": round(x1 / slide_width * 1000),
            "y_max": round(y1 / slide_height * 1000),
        },
    }


def _ole_preview_image(shape) -> tuple[bytes, str] | None:
    """The fallback preview picture every embedded/linked OLE object carries
    (`<p:oleObj>`'s `<mc:Fallback>` branch: `<p:pic><p:blipFill><a:blip>`) —
    this is what actually renders on the slide (confirmed on real documents:
    with `showAsIcon="1"`, it's a generic file-type icon, not a data
    preview). python-pptx exposes no picture API for this shape type
    (`shape.image` doesn't exist on it), so the XML has to be read directly.
    Returns `(raw bytes, extension without the dot)`, or `None` if no
    fallback picture is present."""
    blip = shape._element.find(f".//{qn('p:pic')}/{qn('p:blipFill')}/{qn('a:blip')}")
    if blip is None:
        return None
    rid = blip.get(qn("r:embed"))
    if not rid:
        return None
    try:
        part = shape.part.related_part(rid)
    except KeyError:
        return None
    return part.blob, part.partname.ext.lstrip(".")


def _graphic_data_uri(shape) -> str | None:
    """The `uri` of a `<p:graphicFrame>`'s `<a:graphicData>` — identifies
    what kind of content it holds (chart/table/OLE/SmartArt/...). `None`
    for a non-graphicFrame shape, or a graphicFrame missing that element
    entirely (`find` is safe to call on any shape's element)."""
    graphic_data = shape._element.find(f"{qn('a:graphic')}/{qn('a:graphicData')}")
    return graphic_data.get("uri") if graphic_data is not None else None


def _smartart_text(shape) -> str | None:
    """SmartArt has no rendering support anywhere in this repo either (the
    diagram layout algorithm isn't reproduced by `capture.pptx_capture`'s
    slide renderer, same as PowerPoint's own SmartArt-editing UI needs its
    own dedicated layout engine) — so there's no picture to fall back to
    the way `_ole_preview_image` does for OLE objects. But every text label
    a person actually typed into a SmartArt box is stored, verbatim, as
    plain paragraph runs inside the diagram's data-model part
    (`<dgm:pt><dgm:t>...<a:t>text</a:t>...`), addressable from the shape via
    `<dgm:relIds r:dm="...">`'s data-model relationship — recovering that
    text is far more useful to a text-based/GraphRAG consumer than a
    picture would be anyway. Returns `None` if no text could be recovered
    (e.g. an empty diagram, or a relationship/part that doesn't resolve)."""
    rel_ids = shape._element.find(f".//{_DIAGRAM_NS}relIds")
    if rel_ids is None:
        return None
    dm_rid = rel_ids.get(qn("r:dm"))
    if not dm_rid:
        return None
    try:
        data_part = shape.part.related_part(dm_rid)
    except KeyError:
        return None
    try:
        root = etree.fromstring(data_part.blob)
    except etree.XMLSyntaxError:
        return None
    texts = []
    for pt in root.iter(f"{_DIAGRAM_NS}pt"):
        # type="node" (or unset, its default) is an actual diagram box; other
        # types ("doc", "parTrans", "sibTrans", ...) are structural
        # bookkeeping points with no user-facing text of their own.
        if pt.get("type") not in (None, "node"):
            continue
        text = "".join(run.text or "" for run in pt.findall(f".//{qn('a:t')}")).strip()
        if text:
            texts.append(text)
    return "\n".join(texts) if texts else None


def _format_chart_value(value: float | None) -> str:
    if value is None:
        return ""
    if value == int(value):
        return str(int(value))
    return str(round(value, 4))


def _chart_grid(chart) -> list[list[str]]:
    """A chart is fundamentally the categories/series data it plots — that
    data is exact and machine-readable straight from the chart XML, unlike
    a re-rendered picture (which this repo has no chart-drawing code for
    anyway), so represent it as a Table grid instead: a header row of
    series names, then one row per category. For a chart with multiple
    plots (e.g. a line plot layered over a bar plot), `chart.series`
    already flattens every plot's series into one ordered sequence, and
    categories are taken from the first plot — a reasonable approximation
    since combo charts share one category axis in the overwhelming common
    case. A scatter/bubble chart has no `<c:cat>` categories at all, so it
    degenerates to a header-only grid — not fixed here, since that needs a
    genuinely different (x, y) grid shape, not a missing-categories tweak."""
    categories = [str(c) for c in chart.plots[0].categories] if chart.plots else []
    series = list(chart.series)
    grid = [[""] + [s.name for s in series]]
    for i, category in enumerate(categories):
        row = [category]
        for s in series:
            values = s.values
            row.append(_format_chart_value(values[i] if i < len(values) else None))
        grid.append(row)
    return grid


def _connector_endpoint_ids(shape) -> tuple[int | None, int | None]:
    """The shape ids a connector is glued to, from `<a:stCxn>`/`<a:endCxn>`
    under `<p:nvCxnSpPr>/<p:cNvCxnSpPr>` — present only when the author
    dragged the connector's end onto another shape's connection point in
    PowerPoint's UI. `shape.shape_id` (used to resolve these) is the same
    `<p:cNvPr id>` python-pptx already exposes, unique per slide even across
    nested groups, so no extra XML lookup is needed on the target side."""
    props = shape._element.find(f"{qn('p:nvCxnSpPr')}/{qn('p:cNvCxnSpPr')}")
    if props is None:
        return None, None
    st, end = props.find(qn("a:stCxn")), props.find(qn("a:endCxn"))
    start_id = int(st.get("id")) if st is not None else None
    end_id = int(end.get("id")) if end is not None else None
    return start_id, end_id


def _smallest_node_at_point(nodes: list[Node], x: float, y: float, tolerance: float = 3.0) -> Node | None:
    """The smallest-area content node whose bbox contains (x, y) (0-1000
    slide-space units) — used to infer what an unglued connector visually
    touches. `tolerance` forgives a line landing just short of a shape's
    edge, which is common since the two are independently drawn, not
    snapped. Smallest-area wins so a label sitting on top of a larger
    background shape is preferred over the background."""
    best: Node | None = None
    best_area: float | None = None
    for n in nodes:
        box = n.properties.get("bbox")
        if box is None:
            continue
        if not (box["x_min"] - tolerance <= x <= box["x_max"] + tolerance
                and box["y_min"] - tolerance <= y <= box["y_max"] + tolerance):
            continue
        area = (box["x_max"] - box["x_min"]) * (box["y_max"] - box["y_min"])
        if best_area is None or area < best_area:
            best, best_area = n, area
    return best


def _connector_reference_edge(
    shape, shape_id_to_node: dict[int, Node], content: list[Node],
    slide_width: int, slide_height: int,
) -> Edge | None:
    """A connector (arrow/line) shape draws a relationship, not content — it
    never becomes a Node itself, but what it visually joins is exactly what
    `EdgeType.REFERENCES` means (confirmed on a real document: a "parts
    supplier" slide where a product photo connects, via `stCxn`/`endCxn`-glued
    arrows, to labeled component photos — without this, that relationship
    existed only as pixels, invisible to the graph).

    Each end is resolved independently: if `stCxn`/`endCxn` names a shape, use
    it exactly (and only that — an unresolvable glued target is left
    unresolved rather than guessed from position, since the author's explicit
    intent shouldn't be second-guessed). Absent glue metadata (a plain,
    unattached line), fall back to whichever content node's bbox contains
    that end's actual slide-space point — `_shape_slide_corners(shape)`'s
    first/third corner, the same diagonal `pptx_capture.py`'s `_draw_line`
    renders, so the graph and the pixels the VLM sees can't disagree about
    where the line goes.
    """
    start_id, end_id = _connector_endpoint_ids(shape)
    corners: list[tuple[float, float]] | None = None

    def resolve(shape_id: int | None, corner_index: int) -> Node | None:
        nonlocal corners
        if shape_id is not None:
            return shape_id_to_node.get(shape_id)
        if corners is None:
            corners = _shape_slide_corners(shape)
        x = corners[corner_index][0] / slide_width * 1000
        y = corners[corner_index][1] / slide_height * 1000
        return _smallest_node_at_point(content, x, y)

    start_node = resolve(start_id, 0)
    end_node = resolve(end_id, 2)
    if start_node is None or end_node is None or start_node.id == end_node.id:
        return None
    return Edge(type=EdgeType.REFERENCES, source_id=start_node.id, target_id=end_node.id)


def _classify_shape(
    shape, path_name: str, slide_idx: int, counters: dict, capture_root: Path,
    slide_width: int, slide_height: int,
) -> Node | None:
    layout = _shape_layout_properties(shape, slide_idx, slide_width, slide_height)
    try:
        if shape.has_table:
            table = shape.table
            grid = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            counters["tbl"] += 1
            return Node(
                id=f"content:{path_name}:s{slide_idx}:tbl{counters['tbl']}",
                type=NodeType.TABLE,
                name=f"Table ({len(grid)} rows x {len(grid[0]) if grid else 0} cols)",
                properties={**layout, "grid": grid},
            )
    except (AttributeError, ValueError):
        pass

    if shape.shape_type == MSO_SHAPE_TYPE.CHART:
        chart = shape.chart
        grid = _chart_grid(chart)
        counters["tbl"] += 1
        props = {**layout, "shape_name": shape.name}
        try:
            props["chart_type"] = chart.chart_type.name
        except (AttributeError, ValueError, NotImplementedError):
            pass  # e.g. a combo chart whose plots mix incompatible types
        title = None
        if chart.has_title:
            try:
                title = chart.chart_title.text_frame.text.strip() or None
            except (AttributeError, ValueError):
                pass
        if title:
            props["chart_title"] = title
        return Node(
            id=f"content:{path_name}:s{slide_idx}:tbl{counters['tbl']}",
            type=NodeType.TABLE,
            name=title or f"Chart ({len(grid)} rows x {len(grid[0]) if grid else 0} cols)",
            properties={**props, "grid": grid},
        )

    if _graphic_data_uri(shape) == _SMARTART_GRAPHIC_DATA_URI:
        text = _smartart_text(shape)
        if text:
            counters["txt"] += 1
            return Node(
                id=f"content:{path_name}:s{slide_idx}:txt{counters['txt']}",
                type=NodeType.TEXT,
                name=text[:40] + ("…" if len(text) > 40 else ""),
                properties={**layout, "text": text, "shape_kind": "smartart"},
            )
        return None  # an empty diagram — nothing recoverable, same as an empty text box

    if shape.shape_type in _OLE_SHAPE_TYPES:
        counters["img"] += 1
        props = {**layout, "shape_name": shape.name, "embed_kind": "ole_object"}
        try:
            props["ole_prog_id"] = shape.ole_format.prog_id
        except (AttributeError, ValueError):
            pass  # e.g. a linked (not embedded) object may not expose prog_id
        preview = _ole_preview_image(shape)
        if preview is not None:
            blob, ext = preview
            # Preserve the source preview bytes, including vector formats.
            # Rendering support belongs to the consumer, not native extraction.
            stem = f"{Path(path_name).stem}__s{slide_idx}__img{counters['img']}"
            saved = save_image_bytes(capture_root, stem, blob, ext)
            props["image_path"] = str(saved.resolve())
        return Node(
            id=f"content:{path_name}:s{slide_idx}:img{counters['img']}",
            type=NodeType.IMAGE,
            name=f"Embedded object {counters['img']}",
            properties=props,
        )

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        counters["img"] += 1
        props = {**layout, "shape_name": shape.name}
        # shape.image gives the raw pixels (blob) directly — no separate XML
        # parsing needed (easier than docx). Saved for a pixel-based
        # post-process like a VLM caption (see docx.py's same decision).
        try:
            image = shape.image
            stem = f"{Path(path_name).stem}__s{slide_idx}__img{counters['img']}"
            saved = save_image_bytes(capture_root, stem, image.blob, image.ext)
            props["image_path"] = str(saved.resolve())
        except (AttributeError, ValueError):
            pass  # a corrupted or unsupported image format — keep at least the metadata
        return Node(
            id=f"content:{path_name}:s{slide_idx}:img{counters['img']}",
            type=NodeType.IMAGE,
            name=f"Image {counters['img']}",
            properties=props,
        )

    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            counters["txt"] += 1
            return Node(
                id=f"content:{path_name}:s{slide_idx}:txt{counters['txt']}",
                type=NodeType.TEXT,
                name=text[:40] + ("…" if len(text) > 40 else ""),
                properties={**layout, "text": text},
            )
    return None


def _bbox_overlap(a: dict, b: dict) -> bool:
    return a["x_min"] < b["x_max"] and b["x_min"] < a["x_max"] and a["y_min"] < b["y_max"] and b["y_min"] < a["y_max"]


def _native_pixel_size(shape) -> tuple[int, int] | None:
    """The picture's true source resolution — `shape.image.size` for a
    normal picture, or the OLE fallback preview's own pixel size for an OLE
    object (same fallback `_ole_preview_image` reads). Used only to pick a
    slide-render resolution that doesn't visibly soften a picture when
    annotations are composited onto it (`_annotation_render_dimension`)."""
    try:
        return shape.image.size
    except (AttributeError, ValueError):
        pass
    preview = _ole_preview_image(shape)
    if preview is None:
        return None
    try:
        with PILImage.open(BytesIO(preview[0])) as im:
            return im.size
    except OSError:
        return None


def _annotation_render_dimension(
    image_nodes: list[Node], node_id_to_shape: dict[str, object], slide_width: int, slide_height: int,
    default: int = 1600, hard_max: int = 4096,
) -> int:
    """`PptxSlideRenderer`'s `max_dimension` downsamples the whole slide to
    that size on its long edge — fine for VLM-caption thumbnails, but
    visibly softer than a picture's own resolution when we crop just that
    picture back out of the render (`_render_slide_crops`). Picks the
    smallest `max_dimension` (capped at `hard_max`, the renderer's own
    limit) that still gives every affected picture at least its native
    pixel width after cropping. `image_nodes` is only the pictures being
    annotated — tables have no "native resolution" of their own, so they
    just get whatever this resolves to (at least `default`)."""
    needed = default
    for node in image_nodes:
        shape = node_id_to_shape.get(node.id)
        if shape is None or not shape.width:
            continue
        native = _native_pixel_size(shape)
        if native is None:
            continue
        native_width, _ = native
        # renderer.scale = max_dimension / max(slide_width, slide_height);
        # a picture's crop width in px = scale * shape.width. Solved for
        # max_dimension so that crop width >= native_width.
        needed = max(needed, math.ceil(native_width * max(slide_width, slide_height) / shape.width))
    return min(hard_max, needed)


def _render_slide_crops(
    content: list[Node], annotation_shapes: list, node_id_to_shape: dict[str, object],
    path: Path, slide_idx: int, slide_width: int, slide_height: int, capture_root: Path,
) -> None:
    """Two unrelated gaps share the same fix — crop a region out of one full
    slide render — because `PptxSlideRenderer` already paints every shape
    (pictures, tables, text, connectors) in correct z-order/rotation, so
    reusing its output is cheaper and more accurate than drawing either case
    a second time from scratch:

    1. A shape with no text of its own, drawn directly on top of a picture
       (a highlight box, an unglued arrow) — `_classify_shape` drops it
       (its trailing `return None`), and the picture's saved bytes
       (`shape.image.blob`) are the original embedded file, not a slide
       render, so they never contained the annotation either. Same failure
       mode `capture.xlsx_capture.annotate_standalone_image` exists to fix
       for XLSX's highlight-box-on-photo case.
    2. A native table — unlike XLSX (whose tables get a full visual
       capture), PPTX table extraction had no picture at all: cell shading,
       merged-cell shape, and any picture used as cell fill lived nowhere
       but `grid`'s plain cell text. `PptxSlideRenderer`'s `_draw_table`
       already draws all of that for slide reconstruction — reuse it here
       as `capture_path`, the same property XLSX tables carry. A chart is
       also a Table node (`_chart_grid`) but already carries exact data and
       has no picture-drawing support in the renderer either, so it's
       excluded — only a genuine table shape (`shape.has_table`) qualifies.

    Mutates the affected nodes' properties in place. Renders the slide at
    most once, and only when at least one picture or table actually needs
    it — a slide with neither (the common case for annotations, since most
    pictures have none) costs nothing extra.
    """
    image_nodes = [n for n in content if n.type == NodeType.IMAGE and n.properties.get("image_path")]
    overlapping: dict[str, list] = {}
    for shape in annotation_shapes:
        box = _shape_layout_properties(shape, slide_idx, slide_width, slide_height)["bbox"]
        for node in image_nodes:
            if _bbox_overlap(box, node.properties["bbox"]):
                overlapping.setdefault(node.id, []).append(shape)
    affected_images = [n for n in image_nodes if n.id in overlapping]

    table_nodes = [
        n for n in content
        if n.type == NodeType.TABLE and getattr(node_id_to_shape.get(n.id), "has_table", False)
    ]

    if not affected_images and not table_nodes:
        return

    # Deferred import: capture.pptx_capture imports shape-geometry helpers
    # from this module at its own top level, so importing it back here at
    # module scope would be circular. Both modules are fully initialized by
    # the time extract() actually runs, regardless of which one a caller
    # imports first, so a call-time import is safe.
    from ..capture.pptx_capture import PptxSlideRenderer

    max_dimension = _annotation_render_dimension(affected_images, node_id_to_shape, slide_width, slide_height)
    renderer = PptxSlideRenderer(path, max_dimension=max_dimension)
    slide_png = renderer.render(slide_idx).png
    size = renderer.size
    with PILImage.open(BytesIO(slide_png)) as rendered:
        rendered = rendered.convert("RGB")

        def crop_png(node: Node) -> bytes | None:
            box = node.properties["bbox"]
            crop = (round(box["x_min"] / 1000 * size[0]), round(box["y_min"] / 1000 * size[1]),
                    round(box["x_max"] / 1000 * size[0]), round(box["y_max"] / 1000 * size[1]))
            if crop[2] <= crop[0] or crop[3] <= crop[1]:
                return None  # a degenerate (zero-area) bbox — leave the node untouched
            out = BytesIO()
            rendered.crop(crop).save(out, format="PNG")
            return out.getvalue()

        for node in affected_images:
            png_bytes = crop_png(node)
            if png_bytes is None:
                continue
            stem = f"{Path(node.properties['image_path']).stem}__annotated"
            saved = save_image_bytes(capture_root, stem, png_bytes, "png")
            node.properties["image_path"] = str(saved.resolve())
            node.properties["annotation_shape_count"] = len(overlapping[node.id])

        for node in table_nodes:
            png_bytes = crop_png(node)
            if png_bytes is None:
                continue
            tbl_label = node.id.rsplit(":", 1)[-1]
            stem = f"{path.stem}__s{slide_idx}__{tbl_label}__capture"
            saved = save_image_bytes(capture_root, stem, png_bytes, "png")
            node.properties["capture_path"] = str(saved.resolve())


def extract(path: Path, capture_dir: Path | None = None) -> ArticDocument:
    """`capture_dir`: the directory to save picture shape originals into
    (default: `captures/` next to `path`) — the same convention as
    `docx.extract`."""
    capture_root = resolve_capture_dir(path, capture_dir)
    prs = Presentation(str(path))
    file_n = file_node(path)
    nodes: list[Node] = [file_n]
    edges: list[Edge] = []
    artifacts: list[Node] = []

    for slide_idx, slide in enumerate(prs.slides):
        artifact = Node(
            id=f"artifact:{path.name}:slide{slide_idx}",
            type=NodeType.ARTIFACT,
            name=f"Slide {slide_idx + 1}",
            properties={
                "slide_index": slide_idx,
                "kind": "slide",
                "slide_width": prs.slide_width,
                "slide_height": prs.slide_height,
            },
        )
        nodes.append(artifact)
        artifacts.append(artifact)
        edges.append(Edge(type=EdgeType.PARENT_OF, source_id=file_n.id, target_id=artifact.id))

        counters = {"txt": 0, "tbl": 0, "img": 0}
        flat_shapes = _flatten_shapes(slide.shapes)
        content: list[Node] = []
        shape_id_to_node: dict[int, Node] = {}
        node_id_to_shape: dict[str, object] = {}
        annotation_shapes: list = []
        for shape in flat_shapes:
            n = _classify_shape(
                shape, path.name, slide_idx, counters, capture_root,
                prs.slide_width, prs.slide_height,
            )
            if n is not None:
                content.append(n)
                shape_id_to_node[shape.shape_id] = n
                node_id_to_shape[n.id] = shape
            else:
                # Not table/OLE/picture and no text of its own — normally
                # invisible to the graph. Kept aside so a highlight box (or
                # unglued arrow) drawn directly on top of a picture can still
                # be composited onto it below, instead of silently vanishing.
                annotation_shapes.append(shape)
        content.sort(key=lambda n: (n.properties["bbox"]["y_min"], n.properties["bbox"]["x_min"]))

        _render_slide_crops(
            content, annotation_shapes, node_id_to_shape, path, slide_idx,
            prs.slide_width, prs.slide_height, capture_root,
        )

        nodes.extend(content)
        edges.extend(parent_edges(artifact, content))

        seen_references: set[tuple[str, str]] = set()
        for shape in flat_shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.LINE:
                continue
            edge = _connector_reference_edge(shape, shape_id_to_node, content, prs.slide_width, prs.slide_height)
            if edge is None:
                continue
            # A multi-segment bracket/caliper (e.g. two tick marks + a bar)
            # is several LINE shapes for one conceptual reference — confirmed
            # on a real document (three separate "직선 연결선" shapes all
            # geometrically resolving to the same image/text pair). Keep the
            # graph at one edge per (source, target), not one per segment.
            key = (edge.source_id, edge.target_id)
            if key in seen_references:
                continue
            seen_references.add(key)
            edges.append(edge)

        edges.extend(caption_prefix_edges(content))

    # NEXT between slides (between Artifacts) — the File -> Artifact PARENT_OF edge was already added above
    for a, b in zip(artifacts, artifacts[1:]):
        edges.append(Edge(type=EdgeType.NEXT, source_id=a.id, target_id=b.id))

    return ArticDocument(source_path=str(path.resolve()), format="pptx", nodes=nodes, edges=edges)
