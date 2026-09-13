"""DOCX -> ArticDocument, a fully native path (python-docx only, no LibreOffice/VLM needed).

Walks `document.element.body` in raw XML order so paragraphs (w:p) and
tables (w:tbl) are processed in the order they actually appear — python-docx's
`document.paragraphs`/`document.tables` are each iterated separately and lose
the original document order, so the body is walked directly instead.

Node: Text (paragraph), Table (cell grid; a picture inside one of its cells
      is absorbed as `embedded_images` metadata rather than a separate node
      — the same "Option A: absorb" rule xlsx.py applies to images inside a
      table range, minus the composited capture PNG, since DOCX table
      column widths/row heights are frequently unset and python-docx never
      computes them, so there's no reliable pixel grid to paste onto),
      Image (an inline **raster** image outside any table. A vector shape
      like an arrow/textbox (`w:drawing` present but no `a:blip`) never
      becomes an Image — judging by `w:drawing` presence alone would
      over-include vector shapes too)
Edge: PARENT_OF (Artifact/Table -> content), CAPTION_OF (a heuristic
      proposal — a Text starting with "표 "/"그림 "/"Table "/"Figure " is
      proposed as describing the immediately preceding/following Table/
      Image. A proposal for human review, not a final answer). NEXT is an
      edge only used for order between Artifacts (slides/sheets), and since
      DOCX has only one Artifact ("the whole document"), NEXT is never
      created at all.
"""
from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.xmlchemy import BaseOxmlElement

from ..schema import ArticDocument, Edge, EdgeType, Node, NodeType
from ..scaffold import caption_prefix_edges, file_node, parent_edges, reference_label_edges, resolve_capture_dir, save_image_bytes


def _paragraph_has_image(p: Paragraph) -> bool:
    """True only for a paragraph that has an actual raster image (`a:blip`).

    Judging by `w:drawing` presence alone would catch not just photos but
    vector shapes like arrows/textboxes (`wsp`/`txbxContent`) too. `a:blip`
    is the same tag `_paragraph_image_bytes` looks at when pulling out image
    bytes, so the test criterion and what can actually be saved always
    agree.
    """
    return bool(p._p.findall(".//" + qn("a:blip")))


def _paragraph_image_bytes(p: Paragraph) -> tuple[bytes, str] | None:
    """The raw bytes + extension of the first image in a paragraph. If a
    paragraph has multiple images (rare), only the first is used — matching
    how the current schema only creates one Image node per paragraph (see
    where `_paragraph_has_image` is used).
    """
    blips = p._p.findall(f".//{qn('a:blip')}")
    if not blips:
        return None
    rid = blips[0].get(qn("r:embed"))
    if not rid:
        return None
    part = p.part.related_parts.get(rid)
    if part is None:
        return None
    return part.blob, Path(part.partname).suffix or ".png"



def _body_blocks(element: BaseOxmlElement) -> Iterator[BaseOxmlElement]:
    """Content controls can contain the entire cached Word TOC. Do not
    descend into tables here: their cells remain part of the Table node.
    """
    for child in element:
        if child.tag in {qn("w:p"), qn("w:tbl")}:
            yield child
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _body_blocks(content)


def _outline_level(para: Paragraph) -> int | None:
    """Resolve direct formatting, then the style inheritance chain.
    Word's body-text level 9 explicitly cancels an inherited heading level.
    Do not infer hierarchy from localized/custom style names.
    """
    ppr = para._p.find(qn("w:pPr"))
    style = para.style
    seen: set[str] = set()
    while True:
        level = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
        if level is not None:
            value = level.get(qn("w:val"), "")
            return int(value) if value in {str(i) for i in range(9)} else None
        if style is None or style.style_id in seen:
            return None
        seen.add(style.style_id)
        ppr = style.element.find(qn("w:pPr"))
        style = style.base_style


def _numbering_info(para: Paragraph) -> tuple[int, int] | None:
    """Return Word's explicit list identity and indentation level.

    Automatic numbering is stored in ``w:numPr`` rather than in the visible
    paragraph text. Keeping these values lets relation proposal code detect
    list siblings without fabricating displayed numbers.
    """
    ppr = para._p.find(qn("w:pPr"))
    num_pr = ppr.find(qn("w:numPr")) if ppr is not None else None
    if num_pr is None:
        return None
    num_id = num_pr.find(qn("w:numId"))
    ilvl = num_pr.find(qn("w:ilvl"))
    if num_id is None or ilvl is None:
        return None
    try:
        identity, level = int(num_id.get(qn("w:val"))), int(ilvl.get(qn("w:val")))
        # Word reserves numId=0 for removing numbering, not a shared list.
        return (identity, level) if identity > 0 and 0 <= level <= 8 else None
    except (TypeError, ValueError):
        return None


def _toc_paragraph(para: Paragraph, fields: list[str]) -> bool:
    """Track complex TOC fields across paragraphs, including nested fields.
    Only read cached results; extraction does not update Word fields.
    """
    def in_toc() -> bool:
        return any(re.match(r"\s*TOC(?:\s|$)", instruction, re.I) for instruction in fields)

    toc = in_toc()
    for element in para._p.iter():
        if element.tag == qn("w:fldChar"):
            kind = element.get(qn("w:fldCharType"))
            if kind == "begin":
                fields.append("")
            elif kind == "end" and fields:
                fields.pop()
        elif element.tag == qn("w:instrText") and fields:
            fields[-1] += element.text or ""
        elif element.tag == qn("w:fldSimple"):
            toc = toc or bool(re.match(r"\s*TOC(?:\s|$)", element.get(qn("w:instr"), ""), re.I))
        toc = toc or in_toc()
    return toc


def _native_structure(artifact: Node, content: list[Node]) -> list[Edge]:
    """A heading owns following blocks until an equal/higher level begins.
    TOC entries are navigation, not body headings or section content.
    Explicit internal links are references only when their bookmark is unique.
    """
    stack: list[Node] = []
    edges: list[Edge] = []
    bookmarks: dict[str, list[str]] = {}
    for node in content:
        for name in node.properties.get("bookmarks", []):
            bookmarks.setdefault(name, []).append(node.id)
        level = node.properties.get("outline_level")
        toc = node.properties.get("is_toc", False)
        if level is not None and not toc:
            while stack and stack[-1].properties["outline_level"] >= level:
                stack.pop()
        parent = stack[-1] if stack and not toc else artifact
        edge = parent_edges(parent, [node])[0]
        if not toc and (stack or level is not None):
            edge.properties["structural_source"] = "docx_outline"
        edges.append(edge)
        if level is not None and not toc:
            stack.append(node)
    for node in content:
        for target in node.properties.get("internal_links", []):
            matches = bookmarks.get(target, [])
            if len(matches) == 1 and matches[0] != node.id:
                edges.append(Edge(type=EdgeType.REFERENCES, source_id=node.id, target_id=matches[0],
                                  properties={"structural_source": "docx_bookmark", "bookmark": target}))
    return edges

def extract(path: Path, capture_dir: Path | None = None) -> ArticDocument:
    """`capture_dir`: the directory to save inline image originals into
    (default: `captures/` next to `path`). Same purpose as Table's visual
    capture (xlsx) — keeps the raw bytes around for a pixel-based
    post-process like a VLM caption."""
    capture_root = resolve_capture_dir(path, capture_dir)
    d = docx.Document(str(path))
    file_n = file_node(path)
    artifact = Node(
        id=f"artifact:{path.name}:doc",
        type=NodeType.ARTIFACT,
        name=path.stem,
        properties={"kind": "whole_document"},
    )
    nodes: list[Node] = [file_n, artifact]
    edges: list[Edge] = [Edge(type=EdgeType.PARENT_OF, source_id=file_n.id, target_id=artifact.id)]

    content_nodes: list[Node] = []
    body = d.element.body
    counters = {"p": 0, "tbl": 0, "img": 0}

    # walk body's children in XML order, mapping each to its real python-docx object (Paragraph/Table)
    fields: list[str] = []

    for child in _body_blocks(body):
        tag = child.tag
        if tag == qn("w:p"):
            para = Paragraph(child, d)
            is_toc = _toc_paragraph(para, fields)
            # Cached field and hyperlink text belongs to the paragraph too.
            # Exclude nested textbox paragraphs, which are not inline text.
            text = "".join(
                (e.text or "") if e.tag == qn("w:t") else "\t" if e.tag == qn("w:tab") else "\n"
                for e in child.iter()
                if e.tag in {qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")}
                and next(e.iterancestors(qn("w:p")), None) is child
            ).strip()
            if _paragraph_has_image(para):
                counters["img"] += 1
                props = {"order": len(content_nodes), "nearby_text": text[:80]}
                image_bytes = _paragraph_image_bytes(para)
                if image_bytes is not None:
                    data, ext = image_bytes
                    saved = save_image_bytes(capture_root, f"{path.stem}__img{counters['img']}", data, ext)
                    props["image_path"] = str(saved.resolve())
                img_node = Node(
                    id=f"content:{path.name}:img{counters['img']}",
                    type=NodeType.IMAGE,
                    name=f"Inline image {counters['img']}",
                    properties=props,
                )
                content_nodes.append(img_node)
            if text:
                counters["p"] += 1
                text_node = Node(
                    id=f"content:{path.name}:p{counters['p']}",
                    type=NodeType.TEXT,
                    name=text[:40] + ("…" if len(text) > 40 else ""),
                    properties={
                        "order": len(content_nodes),
                        "text": text,
                        "style": para.style.name if para.style else None,
                    },
                )
                level = _outline_level(para)
                if level is not None and not is_toc:
                    text_node.properties["outline_level"] = level
                numbering = _numbering_info(para)
                if numbering is not None and not is_toc:
                    text_node.properties["list_num_id"], text_node.properties["list_level"] = numbering
                if is_toc:
                    text_node.properties["is_toc"] = True
                for key, tag, attr in (("bookmarks", "w:bookmarkStart", "w:name"),
                                       ("internal_links", "w:hyperlink", "w:anchor")):
                    values = list(dict.fromkeys(e.get(qn(attr)) for e in child.iter(qn(tag)) if e.get(qn(attr))))
                    if values:
                        text_node.properties[key] = values
                content_nodes.append(text_node)
        elif tag == qn("w:tbl"):
            table = Table(child, d)
            grid = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            counters["tbl"] += 1
            table_id = f"content:{path.name}:tbl{counters['tbl']}"
            table_node = Node(
                id=table_id,
                type=NodeType.TABLE,
                name=f"Table {counters['tbl']} ({len(grid)} rows x {len(grid[0]) if grid else 0} cols)",
                properties={"order": len(content_nodes), "grid": grid},
            )
            content_nodes.append(table_node)

            # "Option A: absorb" (xlsx.py's rule for images inside a table
            # range, applied here without a composited capture PNG — see the
            # module docstring): don't miss a picture inside a table cell
            # (looking at cell.text alone silently loses it), but don't give
            # it its own node either. A merged cell has the same _tc element
            # appear multiple times in row.cells, so it's only processed
            # once.
            embedded_images: list[dict] = []
            seen_cells = set()
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    if cell._tc in seen_cells:
                        continue
                    seen_cells.add(cell._tc)
                    for cell_para in cell.paragraphs:
                        if not _paragraph_has_image(cell_para):
                            continue
                        counters["img"] += 1
                        entry = {
                            "table_row": ri,
                            "table_col": ci,
                            "nearby_text": cell_para.text.strip()[:80],
                        }
                        cell_image_bytes = _paragraph_image_bytes(cell_para)
                        if cell_image_bytes is not None:
                            data, ext = cell_image_bytes
                            saved = save_image_bytes(capture_root, f"{path.stem}__img{counters['img']}", data, ext)
                            entry["image_path"] = str(saved.resolve())
                        embedded_images.append(entry)
            if embedded_images:
                table_node.properties["embedded_images"] = embedded_images
                table_node.properties["embedded_image_count"] = len(embedded_images)

    nodes.extend(content_nodes)
    edges.extend(_native_structure(artifact, content_nodes))

    caption_edges = caption_prefix_edges(content_nodes)
    edges.extend(caption_edges)
    edges.extend(reference_label_edges(content_nodes, caption_edges))

    return ArticDocument(source_path=str(path.resolve()), format="docx", nodes=nodes, edges=edges)
