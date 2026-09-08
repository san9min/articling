"""Visually composites a table range into an actual image using only
openpyxl+PIL (no LibreOffice needed).

Converts row/column sizes to pixels to build a canvas, draws text honoring
cell formatting (background color/font/alignment/border/merges), and pastes
the actual bytes of embedded images anchored inside that range at their
computed position. Solves the problem where a photo embedded in an Excel
table (e.g. a sample photo, a measurement-device photo) would otherwise
disappear when only a text grid is extracted — pairs with the "Option A:
absorb" design (an image inside a table range never becomes a separate
Image node, only preserved visually in this capture). Called from
extractors/xlsx.py.

Most image anchors are TwoCellAnchor — using only from+ext would let row-
height estimation error accumulate into the image size calculation and leak
into the next row. TwoCellAnchor computes both the from and to corners
directly, eliminating that error.

But not every TwoCellAnchor means "resize to fit the cell" — with
`editAs="oneCell"` (or "absolute"), `to` is just an approximate coordinate
for screen refresh, not the basis for size; the real size lives in the
shape's own transform extent (left unhandled, a tall photo would get
squashed into one table row's height and badly distorted).

Reads cell style (fill/font/alignment/border, theme colors included) to
apply background color/font weight-color/alignment/merges — see each helper
function for details.

`annotate_standalone_image` applies the same shape compositing
(rectangle/ellipse/connector, via `_load_sheet_shapes`) to a standalone
Image that isn't inside any Table range too — fixing how, unlike table
capture, a highlight box/arrow/label text placed on top of a photo outside
any table used to be entirely lost because only the raw bytes were saved.
`build_sheet_grid` computes the sheet-wide coordinate system once and shares
it — since an image anchor and a shape anchor can reference different cell
ranges (e.g. a highlight box starting a cell or two before the photo),
there's no basis to scope things to "inside that table range" the way table
capture does.

The column-width -> pixel conversion (`_chars_to_px`) is proportional to MDW
(Maximum Digit Width, the "0" glyph's width), and hardcoding this as a
constant (7px) that assumes Excel's new-workbook default font (Calibri
11pt) goes wrong whenever the workbook actually uses a different font (the
Korean engineering documents this project mainly handles commonly default
to Dotum/Malgun Gothic) — the error accumulates the more columns an anchor
is offset by (confirmed: a highlight box on a photo visibly off relative to
the photo's width). `_sheet_mdw_px` measures the actual MDW using the
workbook's default font (finding the local font file via `fc-list`,
measuring the "0"-"9" glyph widths with PIL) and falls back to the old
Calibri approximation (`_FALLBACK_MDW_PX`) only when that font isn't
available locally — this raises accuracy within the limits of what's
possible without an actual rendering engine (LibreOffice etc.), but doesn't
guarantee a perfect pixel match. "Then why not just render with
LibreOffice" was in fact considered — rotated/flipped/cropped shapes' and
images' coordinates get thrown off after PDF conversion, worksheets
spanning multiple pages (confirmed: 5 sheets split into 44 pages) break the
coordinate math entirely, and mapping PDF bookmarks back to the original
sheet is unreliable too (see Experiments.md "Why we don't use LibreOffice as
the renderer") — that path would need the same fail-closed handling, except
the failure blast radius grows to the whole document — so PIL reconstruction
(whose failures stay local) was kept instead.
"""
from __future__ import annotations

import io
import math
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl.utils import get_column_letter
from PIL import Image as PILImage, ImageDraw, ImageFont

from .number_format import format_cell_display

EMU_PER_INCH = 914400
DPI = 96
DEFAULT_ROW_HEIGHT_PT = 15.0
DEFAULT_COL_WIDTH_CHARS = 8.43
# Excel's column width ("character count") -> pixel conversion is
# proportional to MDW (Maximum Digit Width, the "0" glyph's width)
# (ECMA-376 §18.3.1.13). 7px is the MDW of Excel's new-workbook default font,
# Calibri 11pt @ 96dpi — when a document actually uses a different font
# (especially Dotum/Malgun Gothic, the common default font for the Korean
# engineering documents this project mainly handles), this approximation is
# wrong, and the error accumulates across columns (confirmed: a highlight
# box on a photo visibly off relative to the photo's width). `_sheet_mdw_px`
# prefers a value measured from the actual font (falling back to a Korean
# substitute font for Korean fonts when unavailable locally, see
# `_measure_mdw_px`), and only falls back to this constant when even that
# isn't available.
_FALLBACK_MDW_PX = 7.0
MAX_CANVAS_PIXELS = 40_000_000  # skip capturing abnormally large ranges (memory protection)
MAX_CELL_TEXT_CHARS = 300  # prevent slow rendering from a pathologically long cell value

_FONT_REGULAR_CANDIDATES = [
    str(Path.home() / "Library/Fonts/NanumGothic-Regular.ttf"),
    "/System/Library/Fonts/Supplemental/AppleGothic.ttf",
    # Linux fallback — these are hardcoded (not resolved via `fc-list`, unlike
    # _find_font_file_for_family) because this is the last-resort glyph
    # renderer, not the MDW-measurement path; DejaVu ships in the common
    # `fonts-dejavu-core` package and covers Latin text passably. Without
    # this, a Linux install with neither NanumGothic nor DejaVu falls all
    # the way through to PIL's tiny fixed-size `ImageFont.load_default()`,
    # which silently produces badly-sized captures (confirmed: CI's text
    # size measurement diverges enough that fit-shrinking never triggers).
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]
_FONT_BOLD_CANDIDATES = [
    str(Path.home() / "Library/Fonts/NanumGothic-Bold.ttf"),
    "/System/Library/Fonts/Supplemental/AppleGothic.ttf",  # substitute the regular file if no bold file exists (faked bold below)
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
]

_DEFAULT_TEXT_RGB = (30, 30, 30)
_GRID_RGB = (220, 220, 220)  # default light-gray gridline used where there's no actual border formatting
_BORDER_WIDTH_PX = {
    "hair": 1, "thin": 1, "dashed": 1, "dotted": 1, "slantDashDot": 1,
    "mediumDashed": 2, "medium": 2, "double": 2,
    "thick": 3,
}
# The clrScheme entry order that a theme color index (the integer that shows
# up in cell.font.color.theme etc.) refers to. It's a well-known OOXML trap
# that the XML declaration order (dk1,lt1,dk2,lt2,...) and the actual
# reference order swap just the first two — this order is the standard one
# verified across several open-source implementations (SheetJS,
# PhpSpreadsheet, etc.).
_THEME_ORDER = [
    "lt1", "dk1", "lt2", "dk2",
    "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
    "hlink", "folHlink",
]
_THEME_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_XDR_NS = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"
# The names a:schemeClr uses are aliases different from the clrScheme
# declarations (dk1/lt1/dk2/lt2/...) — tx1/bg1/tx2/bg2, meaning "current
# text/background," actually refer to dk1/lt1/dk2/lt2 (a well-known
# indirection in OOXML).
_SCHEME_CLR_ALIASES = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
_DEFAULT_SHAPE_LINE_PX = 1

_font_cache: dict[tuple[str, int], ImageFont.FreeTypeFont] = {}


def _load_capture_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    key = ("bold" if bold else "regular", size)
    if key in _font_cache:
        return _font_cache[key]
    candidates = _FONT_BOLD_CANDIDATES if bold else _FONT_REGULAR_CANDIDATES
    font = None
    for candidate in candidates:
        try:
            font = ImageFont.truetype(candidate, size)
            break
        except OSError:
            continue
    if font is None:
        font = ImageFont.load_default()
    _font_cache[key] = font
    return font


def _pt_to_px(pt: float) -> int:
    return round(pt * DPI / 72)


def _chars_to_px(chars: float, mdw_px: float = _FALLBACK_MDW_PX) -> int:
    """Excel's "character count" column width to pixels. The `+5` is an
    approximation of cell inner padding — strictly this is slightly
    proportional to MDW too (the ECMA-376 formula even has a correction term
    like `Truncate(128/MDW)`), but a constant is plenty in practice for this
    project's scope."""
    return round(chars * mdw_px + 5)


_mdw_cache: dict[tuple[str, float], float | None] = {}
_font_file_cache: dict[str, Path | None] = {}


def _find_font_file_for_family(name: str) -> Path | None:
    """Finds a file among the fonts installed on the system (via `fc-list`
    from fontconfig, if present) whose family name matches this one exactly,
    case-insensitively. `fc-match` isn't used — it silently substitutes some
    other font when it can't find the requested family (confirmed: asking
    for "Dotum" gets substituted with "Verdana"), and an MDW measured from
    the wrong font can be even more wrong than the Calibri approximation.
    Returns None if `fc-list` itself is unavailable (fontconfig not
    installed) or there's no exact match — the caller falls back to
    `_FALLBACK_MDW_PX`.

    Put the other way around: if you install the font the workbook declares
    (often Dotum/Malgun Gothic) into the OS's usual font install path
    (macOS: install via Font Book -> `~/Library/Fonts`, a path fontconfig
    already scans), this function finds and uses that font **with no code
    change** — see Experiments.md's "Install the workbook's font locally to
    improve accuracy," which follows "Why we don't use LibreOffice as the
    renderer." That said, it's an opt-in improvement that only applies on
    the machine that font is installed on — it's a Microsoft-licensed font
    that can't be bundled with the package, so CI/other machines still fall
    back."""
    if name in _font_file_cache:
        return _font_file_cache[name]
    result: Path | None = None
    try:
        out = subprocess.run(
            ["fc-list", "--format=%{file}\t%{family}\n"],
            capture_output=True, text=True, timeout=5,
        )
        if out.returncode == 0:
            target = name.strip().lower()
            for line in out.stdout.splitlines():
                try:
                    file_path, families = line.split("\t", 1)
                except ValueError:
                    continue
                if any(target == fam.strip().lower() for fam in families.split(",")):
                    result = Path(file_path)
                    break
    except (OSError, subprocess.SubprocessError):
        pass  # fontconfig unavailable — fall back
    _font_file_cache[name] = result
    return result


def _measure_mdw_px(font_name: str, size_pt: float) -> float | None:
    """Measures the widest glyph among "0"-"9" (the MDW) in pixels, using the
    `font_name`/`size_pt` font. Returns None if no font by that exact name
    is available locally (the caller falls back to `_FALLBACK_MDW_PX`) —
    measuring instead with the Korean substitute font this module uses for
    label-text rendering (`_FONT_REGULAR_CANDIDATES`, NanumGothic/
    AppleGothic) was also tried (confirmed on a screw fastening-strength
    review document, substituting Dotum -> NanumGothic), but the substitute
    font's actual MDW came out even further from the Calibri approximation
    than the original font's would be (overshoot 22% -> 53%), so it was
    dropped — a case showing that an unverified font-substitution guess can
    be worse than not guessing at all. Cached per (font_name, size_pt)
    combination — a sheet with dozens to hundreds of columns still only gets
    measured once.

    Additional finding (2026-09-05): also tried Baekmuk Dotum, a free-license
    font billed as "metrically similar to the Dotum/Gulim family" (a
    SourceForge distribution, with the font name table rewritten to "Dotum"
    so `_find_font_file_for_family` would find it). The measured MDW came
    out to 8.0 instead of 7.0 (the Calibri fallback), making actual shape
    position error worse, not better (the same document's relative shape
    positions drifted from 78-122% to 87-138%) — reconfirming that a
    "visually similar substitute font" is in no way guaranteed to match even
    character-width (advance width) metrics. So both NanumGothic and Baekmuk
    are dropped — of every substitute font tried so far, none beats the
    Calibri fallback (7.0)."""
    key = (font_name, size_pt)
    if key in _mdw_cache:
        return _mdw_cache[key]
    font_path = _find_font_file_for_family(font_name)
    mdw: float | None = None
    if font_path is not None:
        try:
            size_px = round(size_pt * DPI / 72)
            font = ImageFont.truetype(str(font_path), size_px)
            mdw = max(font.getlength(str(d)) for d in range(10))
        except OSError:
            mdw = None
    _mdw_cache[key] = mdw
    return mdw


def _workbook_default_font(wb) -> tuple[str, float]:
    """The workbook's default ("Normal") cell-style font name/size. Cells can
    each use a different font, but Excel itself also computes column
    width->pixel conversion using this one default-style font's MDW
    uniformly (so that a column mixing different fonts still gets one
    width) — we follow the same assumption. Falls back to Excel's
    new-workbook default (Calibri 11) if not found."""
    try:
        styles = wb._named_styles
        if styles:
            font = styles[0].font
            return (font.name or "Calibri"), (font.sz or 11.0)
    except (AttributeError, IndexError):
        pass
    return "Calibri", 11.0


def _sheet_mdw_px(ws) -> float:
    """MDW (in pixels) measured with the default font of the workbook this
    worksheet belongs to. Falls back to `_FALLBACK_MDW_PX` (a Calibri 11
    approximation) if that font isn't available locally."""
    name, size = _workbook_default_font(ws.parent)
    return _measure_mdw_px(name, size) or _FALLBACK_MDW_PX


def sheet_font_metrics_measured(ws) -> tuple[bool, str]:
    """`(True, font_name)` if this worksheet's default font's actual MDW was
    measured; `(False, font_name)` if that font wasn't available locally and
    it fell back to `_FALLBACK_MDW_PX` (a Calibri approximation).

    `extractors/xlsx.py` keeps this as the `column_width_font_measured`
    property on the Artifact (worksheet) node, and warns when it's `False`
    — since `editAs="twoCell"` (genuinely resized to fit the cell)
    shape/image positions depend on the approximated column width, the
    caller needs to know it was computed without the actual measured font
    so the cause is immediately clear when a position looks off relative to
    another anchor on the same photo (see Experiments.md "Install the
    workbook's font locally to improve accuracy" — though not just any
    substitute font will do, see the same section's NanumGothic/Baekmuk
    findings). The font name is returned alongside so the caller can put it
    in the warning message without calling the (internal) function
    `_workbook_default_font` separately again."""
    name, size = _workbook_default_font(ws.parent)
    return _measure_mdw_px(name, size) is not None, name


def _emu_to_px(emu: float) -> int:
    return round(emu / EMU_PER_INCH * DPI)


def _cumulative_offsets(sizes: dict[int, int], start: int, end: int) -> dict[int, int]:
    offsets = {}
    acc = 0
    for i in range(start, end + 1):
        offsets[i] = acc
        acc += sizes[i]
    offsets[end + 1] = acc
    return offsets


def _row_height_px(ws, row: int) -> int:
    dim = ws.row_dimensions.get(row)
    if dim and dim.hidden:
        return 0  # a hidden row is 0px on screen/print, regardless of its declared height
    pt = dim.height if dim and dim.height else (ws.sheet_format.defaultRowHeight or DEFAULT_ROW_HEIGHT_PT)
    return _pt_to_px(pt)


def _sheet_default_col_width_chars(ws) -> float:
    """Gets this sheet's actual default column width (in "character count"
    units). This value applies to every column with no explicit width set
    (the overwhelming majority), so getting it wrong here accumulates error
    across many columns — confirmed (in a screw fastening-strength review
    document) as one cause of a highlight box on a photo visibly off
    relative to the photo's actual width.

    Uses `ws.sheet_format.defaultColWidth` as-is if it's declared. Otherwise
    (as in this document) falls back to `baseColWidth` (the same "character
    count" unit — ECMA-376 §18.3.1.13) — unconditionally using
    `DEFAULT_COL_WIDTH_CHARS` (8.43, a value specific to Excel new-workbook's
    Calibri 11pt default) would ignore that document's actual declaration
    (e.g. `baseColWidth=8`). The MDW itself that `_chars_to_px` uses to turn
    "character count" into pixels varies by font, so it isn't handled here —
    the caller (`_col_width_px`) handles it separately via `_sheet_mdw_px`."""
    fmt = ws.sheet_format
    if fmt.defaultColWidth:
        return fmt.defaultColWidth
    if fmt.baseColWidth:
        return fmt.baseColWidth
    return DEFAULT_COL_WIDTH_CHARS


def _col_width_px(ws, col: int, mdw_px: float = _FALLBACK_MDW_PX) -> int:
    letter = get_column_letter(col)
    dim = ws.column_dimensions.get(letter)
    if dim and dim.hidden:
        return 0  # a hidden column is 0px on screen/print, regardless of its declared width
    chars = dim.width if dim and dim.width else _sheet_default_col_width_chars(ws)
    return _chars_to_px(chars, mdw_px)


def rects_overlap(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> bool:
    """Public function — exposed across the module boundary because
    `extractors/xlsx.py` needs the same overlap rule as the rest of this
    module when it checks whether standalone images outside any table
    overlap each other (e.g. a small icon photo placed separately on top of
    another photo)."""
    ax1, ay1, ax2, ay2 = a
    bx1, by1, bx2, by2 = b
    return ax1 < bx2 and bx1 < ax2 and ay1 < by2 and by1 < ay2


def build_sheet_grid(ws, max_row: int, max_col: int) -> tuple[dict[int, int], dict[int, int], int, int]:
    """Computes the whole sheet's (1..max_row, 1..max_col) cumulative
    row/column pixel offsets once.

    `capture_table_image` only needs the table range (±padding) so it
    computes it locally each time, but `annotate_standalone_image`
    (annotation compositing over a standalone image) needs a coordinate
    system spanning the whole sheet, since an image anchor and a shape
    anchor can reference different cell ranges (e.g. a highlight box
    starting a cell or two before the photo) — called once per sheet and
    reused (`extractors/xlsx.py`) so it isn't recomputed per image. Column
    width is converted to pixels using the MDW measured from the workbook's
    default font (`_sheet_mdw_px`, only when that font is available
    locally) — measured once per sheet, not repeated per column."""
    if max_row <= 0 or max_col <= 0:
        return {}, {}, 0, 0
    mdw_px = _sheet_mdw_px(ws)
    row_h = {r: _row_height_px(ws, r) for r in range(1, max_row + 1)}
    col_w = {c: _col_width_px(ws, c, mdw_px) for c in range(1, max_col + 1)}
    row_off = _cumulative_offsets(row_h, 1, max_row)
    col_off = _cumulative_offsets(col_w, 1, max_col)
    return row_off, col_off, col_off[max_col + 1], row_off[max_row + 1]


def _parse_theme_colors(wb) -> list[str] | None:
    """Returns the workbook theme's 12 colors (clrScheme) as a list of RGB
    hex strings in `_THEME_ORDER`. Returns None if there's no theme or
    parsing fails (theme colors can't be used -> fall back to default
    colors)."""
    raw = getattr(wb, "loaded_theme", None)
    if not raw:
        return None
    try:
        root = ET.fromstring(raw)
        scheme = root.find(f".//{_THEME_NS}clrScheme")
        if scheme is None:
            return None
        found: dict[str, str] = {}
        for child in scheme:
            tag = child.tag.replace(_THEME_NS, "")
            srgb = child.find(f"{_THEME_NS}srgbClr")
            sysclr = child.find(f"{_THEME_NS}sysClr")
            if srgb is not None and srgb.get("val"):
                found[tag] = srgb.get("val")
            elif sysclr is not None and sysclr.get("lastClr"):
                found[tag] = sysclr.get("lastClr")
        return [found.get(name, "000000") for name in _THEME_ORDER]
    except ET.ParseError:
        return None


def _apply_tint(rgb: tuple[int, int, int], tint: float) -> tuple[int, int, int]:
    """Approximates Excel's theme-color tint (-1.0 = fully black to +1.0 =
    fully white) by applying it per channel. Real Excel computes it based on
    HSL luminance, but a per-channel linear approximation is visually close
    enough for table-capture purposes."""
    if not tint:
        return rgb
    if tint < 0:
        factor = 1 + tint
        return tuple(max(0, round(ch * factor)) for ch in rgb)  # type: ignore[return-value]
    return tuple(min(255, round(ch * (1 - tint) + 255 * tint)) for ch in rgb)  # type: ignore[return-value]


def _resolve_color(color, theme_colors: list[str] | None) -> tuple[int, int, int] | None:
    """Resolves an openpyxl `Color` (font.color/fill.fgColor/border
    side.color) to RGB. Returns None if it can't be resolved (a legacy
    indexed palette, no theme info, etc.) — the caller falls back to a
    default color."""
    if color is None:
        return None
    try:
        color_type = color.type
    except AttributeError:
        return None

    if color_type == "rgb" and isinstance(color.rgb, str) and len(color.rgb) >= 6:
        hex_part = color.rgb[-6:]
        try:
            return tuple(int(hex_part[i : i + 2], 16) for i in (0, 2, 4))  # type: ignore[return-value]
        except ValueError:
            return None
    if color_type == "theme" and theme_colors is not None:
        idx = color.theme
        if idx is not None and 0 <= idx < len(theme_colors):
            base_hex = theme_colors[idx]
            try:
                base_rgb = tuple(int(base_hex[i : i + 2], 16) for i in (0, 2, 4))
            except ValueError:
                return None
            return _apply_tint(base_rgb, color.tint or 0.0)
    return None  # e.g. an indexed palette — not supported


def _resolve_xml_color(container: ET.Element, theme_colors: list[str] | None) -> tuple[int, int, int] | None:
    """Resolves an `<a:srgbClr>` or `<a:schemeClr>` child directly under a
    container element like `<a:solidFill>`/`<a:ln>` to RGB. The raw-XML
    version of `_resolve_color` (for openpyxl `Color` objects) — shapes
    (`xdr:sp`) aren't parsed by openpyxl, so the drawing XML has to be dug
    into directly. Fine-grained shading like lumMod/lumOff isn't applied
    (an approximation)."""
    srgb = container.find(f"{_THEME_NS}srgbClr")
    if srgb is not None and srgb.get("val"):
        hexv = srgb.get("val")
        try:
            return tuple(int(hexv[i : i + 2], 16) for i in (0, 2, 4))  # type: ignore[return-value]
        except ValueError:
            return None
    scheme = container.find(f"{_THEME_NS}schemeClr")
    if scheme is not None and theme_colors is not None:
        name = _SCHEME_CLR_ALIASES.get(scheme.get("val") or "", scheme.get("val"))
        if name in _THEME_ORDER:
            base_hex = theme_colors[_THEME_ORDER.index(name)]
            try:
                return tuple(int(base_hex[i : i + 2], 16) for i in (0, 2, 4))  # type: ignore[return-value]
            except ValueError:
                return None
    return None


def _xml_alpha(container: ET.Element) -> float:
    """Converts `<a:alpha val="50000"/>` under a color element
    (`<a:srgbClr>`/`<a:schemeClr>`) to 0.0-1.0 opacity. Fully opaque (1.0) if
    absent."""
    for child in container:
        if child.tag in (f"{_THEME_NS}srgbClr", f"{_THEME_NS}schemeClr"):
            alpha_elem = child.find(f"{_THEME_NS}alpha")
            if alpha_elem is not None and alpha_elem.get("val"):
                try:
                    return int(alpha_elem.get("val")) / 100_000
                except ValueError:
                    return 1.0
    return 1.0


_BENT_OR_CURVED_CONNECTOR_PREFIXES = ("bentConnector", "curvedConnector")


def _xfrm_transform(sppr: ET.Element | None) -> tuple[bool, bool, float]:
    """Reads `spPr/xfrm`'s `flipH`/`flipV`/`rot` (in 1/60000ths of a degree,
    clockwise). All defaults if `xfrm` itself is absent (as for most shapes,
    where the anchor position is the only basis for size)."""
    if sppr is None:
        return False, False, 0.0
    xfrm = sppr.find(f"{_THEME_NS}xfrm")
    if xfrm is None:
        return False, False, 0.0
    flip_h = xfrm.get("flipH") == "1"
    flip_v = xfrm.get("flipV") == "1"
    rot_raw = xfrm.get("rot")
    rot_deg = int(rot_raw) / 60_000.0 if rot_raw else 0.0
    return flip_h, flip_v, rot_deg


def _shape_style(
    sp: ET.Element, theme_colors: list[str] | None
) -> tuple[str, tuple[int, int, int, float] | None, tuple[tuple[int, int, int], int] | None, float]:
    """Extracts (the preset geometry name, fill RGBA, outline (RGB, width
    px), rotation angle (degrees, clockwise)) from an `<xdr:sp>` (shape).
    Can't read a shape with no explicit color in `spPr` (only a theme style
    reference `<xdr:style>` with no override) — in that case both fill and
    outline come back None and the caller doesn't draw that shape (no color
    is invented by guessing). Confirmed in practice that when a person
    changes a shape's formatting (fill/outline) directly in Excel, `spPr`
    gets an explicit override written alongside it — only a shape left at
    the pure theme default hits this limitation. Also usable as-is for
    `<xdr:cxnSp>` (connectors) — only `spPr/ln` (line color/width) is
    meaningful for a connector, and the caller can ignore
    `solidFill`/`prstGeom`."""
    prst = "rect"
    fill: tuple[int, int, int, float] | None = None
    outline: tuple[tuple[int, int, int], int] | None = None

    sppr = sp.find(f"{_XDR_NS}spPr")
    if sppr is None:
        return prst, fill, outline, 0.0

    geom = sppr.find(f"{_THEME_NS}prstGeom")
    if geom is not None and geom.get("prst"):
        prst = geom.get("prst")  # type: ignore[assignment]

    solid = sppr.find(f"{_THEME_NS}solidFill")
    if solid is not None:
        rgb = _resolve_xml_color(solid, theme_colors)
        if rgb is not None:
            fill = (*rgb, _xml_alpha(solid))

    ln = sppr.find(f"{_THEME_NS}ln")
    if ln is not None:
        ln_solid = ln.find(f"{_THEME_NS}solidFill")
        if ln_solid is not None:
            rgb = _resolve_xml_color(ln_solid, theme_colors)
            if rgb is not None:
                width_emu = ln.get("w")
                width_px = _emu_to_px(int(width_emu)) if width_emu else _DEFAULT_SHAPE_LINE_PX
                outline = (rgb, max(1, width_px))

    _, _, rot_deg = _xfrm_transform(sppr)
    return prst, fill, outline, rot_deg


def _has_arrowhead(ln: ET.Element | None, tag: str) -> bool:
    """Whether `<a:headEnd>`/`<a:tailEnd type="triangle|stealth|arrow|diamond|oval"/>`
    exists under `<a:ln>`. No arrowhead if `type="none"` or the element
    itself is absent. Drawn as a triangle approximation regardless of the
    exact shape (triangle/diamond/oval etc.)."""
    if ln is None:
        return False
    end = ln.find(f"{_THEME_NS}{tag}")
    return end is not None and end.get("type") not in (None, "none")


def _load_sheet_shapes(xlsx_path: Path, ws) -> list[tuple[ET.Element, ET.Element, str]]:
    """Returns the top-level anchors of `<xdr:sp>` (filled shapes like
    rectangles/ellipses), `<xdr:cxnSp>` (connectors — arrows/lines), and
    `<xdr:grpSp>` (shapes a person selected and "grouped") from the drawing
    part the worksheet is linked to, as a list of `(anchor, elem, kind)`
    triples (`kind` is `"sp"`/`"cxnSp"`/`"grpSp"`). Actual embedded photos
    (`<xdr:pic>`) aren't included. A `kind == "grpSp"` entry can't be drawn
    as-is — the caller has to unpack the inner sp/cxnSp with
    `_expand_group_shapes(elem, group_rect)` (a caller that only needs the
    group's row/col range, like `shape_anchor_max_row_col`, doesn't need to
    expand the group — the top-level anchor's from/to is enough).

    openpyxl doesn't parse these shapes at all — `find_images()`
    (openpyxl/reader/drawings.py) only scans for `pic` and warns "Shapes and
    drawings will be lost." They never show up via `ws._images`, so the only
    option is to open the original xlsx as a zip directly and re-parse the
    drawing XML."""
    try:
        rel = next(r for r in (getattr(ws, "_rels", None) or []) if r.Type.endswith("/drawing"))
    except StopIteration:
        return []
    target = rel.target.lstrip("/")
    try:
        with zipfile.ZipFile(xlsx_path) as zf:
            raw = zf.read(target)
    except (KeyError, OSError):
        return []
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        return []

    triples: list[tuple[ET.Element, ET.Element, str]] = []
    for anchor_tag in ("twoCellAnchor", "oneCellAnchor"):
        for anchor in root.findall(f"{_XDR_NS}{anchor_tag}"):
            sp = anchor.find(f"{_XDR_NS}sp")
            if sp is not None:
                triples.append((anchor, sp, "sp"))
                continue
            cxn = anchor.find(f"{_XDR_NS}cxnSp")
            if cxn is not None:
                triples.append((anchor, cxn, "cxnSp"))
                continue
            grp = anchor.find(f"{_XDR_NS}grpSp")
            if grp is not None:
                triples.append((anchor, grp, "grpSp"))
    return triples


def _expand_group_shapes(
    grp: ET.Element,
    group_rect: tuple[int, int, int, int],
) -> list[tuple[tuple[int, int, int, int], ET.Element, str]]:
    """Unpacks the sp/cxnSp inside an `xdr:grpSp` that a person actually
    grouped by selecting several shapes, returning them as a list of
    (absolute pixel rectangle, element, kind).

    Each child's position inside the group is written not relative to its
    own `spPr/xfrm` (off/ext) but relative to the **group's internal
    coordinate system** (`grpSpPr/xfrm`'s `chOff`/`chExt` — ECMA-376
    `CT_GroupTransform2D`). The group itself is already a top-level anchor
    (`twoCellAnchor`/`oneCellAnchor`), so `group_rect` (the caller gets this
    beforehand via `_xml_anchor_rect`) is its absolute pixel rectangle, and
    linearly mapping (translate+scale only, no rotation) each child's
    relative position (as a 0-1 fraction) within the `chOff`~`chOff+chExt`
    range onto that `group_rect` gives the child's absolute position — since
    the group coordinate system itself expresses "what % within this
    rectangle," this is exact arithmetic when there's no rotation, not a
    guess.

    If the group itself is rotated (`rot`), the internal layout would need
    that rotation composited back in to be accurate, which is out of this
    project's scope (the same "don't guess" principle as skipping
    `bentConnector`/`curvedConnector`) — such a group is skipped entirely.
    Nested groups (a group inside a group) and an `xdr:pic` (an actual
    image — getting its pixels needs unpacking drawing.xml's relationship
    file too) inside a group are also out of scope for now."""
    grp_sppr = grp.find(f"{_XDR_NS}grpSpPr")
    xfrm = grp_sppr.find(f"{_THEME_NS}xfrm") if grp_sppr is not None else None
    if xfrm is None or xfrm.get("rot") not in (None, "0"):
        return []  # no xfrm, or a rotated group (out of scope) — skip safely
    ch_off = xfrm.find(f"{_THEME_NS}chOff")
    ch_ext = xfrm.find(f"{_THEME_NS}chExt")
    if ch_off is None or ch_ext is None:
        return []  # can't resolve the group's internal coordinate system, so can't resolve child positions either
    ch_x, ch_y = int(ch_off.get("x", "0")), int(ch_off.get("y", "0"))
    ch_cx, ch_cy = int(ch_ext.get("cx", "0")), int(ch_ext.get("cy", "0"))
    if ch_cx == 0 or ch_cy == 0:
        return []

    gx1, gy1, gx2, gy2 = group_rect
    gw, gh = gx2 - gx1, gy2 - gy1

    results: list[tuple[tuple[int, int, int, int], ET.Element, str]] = []
    for tag, kind in ((f"{_XDR_NS}sp", "sp"), (f"{_XDR_NS}cxnSp", "cxnSp")):
        for child in grp.findall(tag):
            child_sppr = child.find(f"{_XDR_NS}spPr")
            child_xfrm = child_sppr.find(f"{_THEME_NS}xfrm") if child_sppr is not None else None
            if child_xfrm is None:
                continue
            off = child_xfrm.find(f"{_THEME_NS}off")
            ext = child_xfrm.find(f"{_THEME_NS}ext")
            if off is None or ext is None:
                continue
            cx, cy = int(off.get("x", "0")), int(off.get("y", "0"))
            ccx, ccy = int(ext.get("cx", "0")), int(ext.get("cy", "0"))
            fx1, fy1 = (cx - ch_x) / ch_cx, (cy - ch_y) / ch_cy
            fx2, fy2 = fx1 + ccx / ch_cx, fy1 + ccy / ch_cy
            rect = (
                round(gx1 + fx1 * gw), round(gy1 + fy1 * gh),
                round(gx1 + fx2 * gw), round(gy1 + fy2 * gh),
            )
            results.append((rect, child, kind))
    return results


def _flatten_shape_entries(
    xlsx_path: Path,
    ws,
    col_off: dict[int, int],
    row_off: dict[int, int],
    canvas_w: int,
    canvas_h: int,
) -> list[tuple[ET.Element, tuple[int, int, int, int], ET.Element, str]]:
    """Flattens the top-level shapes/connectors/groups found by
    `_load_sheet_shapes` into a list of (top-level anchor, for scope checks
    like "is it in the table range" / "does it overlap this image", absolute
    pixel rectangle, element to draw, kind) that can actually be drawn —
    shared preparation used by `capture_table_image`/
    `annotate_standalone_image`. After flattening, kind is always just
    `"sp"`/`"cxnSp"` (groups disappear, replaced by their inner children, see
    `_expand_group_shapes`). A child that came out of a group still uses that
    group's top-level anchor for the caller's scope checks (in the table
    range or not, overlapping a particular image or not) — once something
    belongs to a group, it's natural to treat the whole group as one unit
    for being in/out of scope."""
    entries: list[tuple[ET.Element, tuple[int, int, int, int], ET.Element, str]] = []
    for anchor, elem, kind in _load_sheet_shapes(xlsx_path, ws):
        if kind == "grpSp":
            group_rect = _xml_anchor_rect(anchor, col_off, row_off, canvas_w, canvas_h)
            if group_rect is None:
                continue
            for child_rect, child_elem, child_kind in _expand_group_shapes(elem, group_rect):
                entries.append((anchor, child_rect, child_elem, child_kind))
            continue
        rect = _xml_anchor_rect(anchor, col_off, row_off, canvas_w, canvas_h)
        if rect is None:
            continue
        entries.append((anchor, rect, elem, kind))
    return entries


def shape_anchor_max_row_col(xlsx_path: Path, ws) -> tuple[int, int]:
    """Returns the largest (row, col) (1-based) reached by any shape
    (`xdr:sp`/`xdr:cxnSp`) anchor in the drawing part linked to `ws`.

    Used by `extractors/xlsx.py`'s `_content_bounds` when determining the
    scan range, to include shape-only areas (e.g. a highlight box drawn a
    bit larger than the photo) that value cells/merges/tables alone can't
    catch — otherwise the sheet-wide coordinate system `build_sheet_grid`
    builds wouldn't cover that shape's anchor cell, and
    `annotate_standalone_image` couldn't see the overlap at all."""
    max_row = max_col = 0
    for anchor, _elem, _kind in _load_sheet_shapes(xlsx_path, ws):
        for tag in (f"{_XDR_NS}from", f"{_XDR_NS}to"):
            node = anchor.find(tag)
            if node is None:
                continue
            r = int(node.findtext(f"{_XDR_NS}row", "0")) + 1
            c = int(node.findtext(f"{_XDR_NS}col", "0")) + 1
            max_row, max_col = max(max_row, r), max(max_col, c)
    return max_row, max_col


def _xml_anchor_rect(
    anchor: ET.Element,
    col_off: dict[int, int],
    row_off: dict[int, int],
    canvas_w: int,
    canvas_h: int,
) -> tuple[int, int, int, int] | None:
    """Computes the pixel rectangle of an `<xdr:twoCellAnchor>`/
    `<xdr:oneCellAnchor>` (raw XML). A raw-XML repeat of the same
    from/to/ext + editAs rules `capture_table_image` uses for embedded
    images (can't reuse that one directly since this is an ElementTree, not
    an openpyxl object, so it's rewritten in parallel)."""
    frm = anchor.find(f"{_XDR_NS}from")
    if frm is None:
        return None
    fc = int(frm.findtext(f"{_XDR_NS}col", "0")) + 1
    fco = int(frm.findtext(f"{_XDR_NS}colOff", "0"))
    fr = int(frm.findtext(f"{_XDR_NS}row", "0")) + 1
    fro = int(frm.findtext(f"{_XDR_NS}rowOff", "0"))
    if fc not in col_off or fr not in row_off:
        return None
    x = col_off[fc] + _emu_to_px(fco)
    y = row_off[fr] + _emu_to_px(fro)

    edit_as = anchor.get("editAs")
    to = anchor.find(f"{_XDR_NS}to")
    ext = anchor.find(f"{_XDR_NS}ext")
    fixed_size = to is None or edit_as in ("oneCell", "absolute")

    if to is not None and not fixed_size:
        tc = int(to.findtext(f"{_XDR_NS}col", "0")) + 1
        tco = int(to.findtext(f"{_XDR_NS}colOff", "0"))
        tr = int(to.findtext(f"{_XDR_NS}row", "0")) + 1
        tro = int(to.findtext(f"{_XDR_NS}rowOff", "0"))
        x2 = col_off[tc] + _emu_to_px(tco) if tc in col_off else canvas_w
        y2 = row_off[tr] + _emu_to_px(tro) if tr in row_off else canvas_h
        w, h = max(1, x2 - x), max(1, y2 - y)
    elif ext is not None:
        w = max(1, _emu_to_px(int(ext.get("cx", "0"))))
        h = max(1, _emu_to_px(int(ext.get("cy", "0"))))
    else:
        return None  # no basis for a size (an abnormal case: a oneCellAnchor with no ext either)

    return x, y, x + w, y + h


def _shape_text_lines(
    sp: ET.Element, theme_colors: list[str] | None
) -> list[tuple[str, tuple[int, int, int], int]]:
    """Extracts the paragraphs inside an `<xdr:sp>`'s `<xdr:txBody>` as a
    list of (text, RGB, font pt) tuples. A real annotation label (e.g. "2
    left fixing screws") is almost always one run per paragraph, so if a
    paragraph has multiple runs, only the text is concatenated and
    formatting is approximated from the first run. Falls back to black if
    there's no explicit color — unlike shape fill/outline (never drawn at
    all if the color can't be read, see `_shape_style`), for text it's
    better to at least show it in a default color than not draw it, to
    preserve the label's existence at all."""
    txbody = sp.find(f"{_XDR_NS}txBody")
    if txbody is None:
        return []
    lines: list[tuple[str, tuple[int, int, int], int]] = []
    for p in txbody.findall(f"{_THEME_NS}p"):
        runs = p.findall(f"{_THEME_NS}r")
        text = "".join(r.findtext(f"{_THEME_NS}t") or "" for r in runs)
        if not text.strip():
            continue
        rgb = (0, 0, 0)
        size_pt = 11
        rpr = runs[0].find(f"{_THEME_NS}rPr") if runs else None
        if rpr is not None:
            solid = rpr.find(f"{_THEME_NS}solidFill")
            if solid is not None:
                resolved = _resolve_xml_color(solid, theme_colors)
                if resolved is not None:
                    rgb = resolved
            sz_attr = rpr.get("sz")  # in hundredths of a point (OOXML)
            if sz_attr:
                try:
                    size_pt = max(1, round(int(sz_attr) / 100))
                except ValueError:
                    pass
        lines.append((text, rgb, size_pt))
    return lines


def _draw_shape_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    lines: list[tuple[str, tuple[int, int, int], int]],
) -> None:
    """Draws paragraphs stacked and vertically centered within a shape's
    bbox. Doesn't wrap by character — in practice an annotation label is one
    or two words shorter than the shape's width, so this is enough, and
    pulling in `_wrap_lines` would add more complexity than this
    approximate renderer warrants."""
    x1, y1, x2, y2 = box
    box_w, box_h = max(1, x2 - x1), max(1, y2 - y1)
    rendered = []
    total_h = 0
    for text, rgb, size_pt in lines:
        font_px = min(40, max(8, round(size_pt * DPI / 72 * 0.85)))
        font = _load_capture_font(font_px)
        line_h = round(font_px * 1.3)
        rendered.append((text, rgb, font, line_h))
        total_h += line_h
    y = y1 + max(0, (box_h - total_h) // 2)
    for text, rgb, font, line_h in rendered:
        line_w = draw.textlength(text, font=font)
        x = x1 + max(0, (box_w - line_w) / 2)
        draw.text((x, y), text, fill=rgb, font=font)
        y += line_h


def _paint_shape(
    canvas: PILImage.Image,
    rect: tuple[int, int, int, int],
    prst: str,
    fill: tuple[int, int, int, float] | None,
    outline: tuple[tuple[int, int, int], int] | None,
    rotation_deg: float = 0.0,
    text_lines: list[tuple[str, tuple[int, int, int], int]] | None = None,
) -> None:
    """Draws a shape (rectangle/ellipse) within `rect` (bbox) and composites
    it onto `canvas` (RGB). Always drawn on a separate RGBA tile first and
    then pasted, uniformly — both semi-transparent fills and rotation go
    through the same path (there's no separate special case for drawing
    directly on the canvas), so branching doesn't grow. Rotation pivots on
    the bbox center (the same convention as OOXML `rot`). If `text_lines` is
    present (`_shape_text_lines`), the label text on the shape is drawn on
    the same tile along with the rotation — it's visually correct for the
    text on a rotated shape to rotate along with it.
    """
    x1, y1, x2, y2 = rect
    w, h = max(1, x2 - x1), max(1, y2 - y1)
    tile = PILImage.new("RGBA", (w, h), (0, 0, 0, 0))
    tile_draw = ImageDraw.Draw(tile)
    shape_fn = tile_draw.ellipse if prst == "ellipse" else tile_draw.rectangle
    local_rect = (0, 0, w - 1, h - 1)
    if fill is not None:
        r, g, b, alpha = fill
        shape_fn(local_rect, fill=(r, g, b, round(alpha * 255)))
    if outline is not None:
        (r, g, b), width = outline
        shape_fn(local_rect, outline=(r, g, b), width=width)
    if text_lines:
        _draw_shape_text(tile_draw, local_rect, text_lines)

    if rotation_deg:
        # OOXML `rot` is clockwise, while PIL Image.rotate() turns a
        # positive angle counterclockwise, so the sign must be flipped to
        # match visually. expand=True grows the canvas to fit the larger
        # bbox after rotation, and it's pasted back centered on the
        # original center point.
        tile = tile.rotate(-rotation_deg, expand=True, resample=PILImage.BICUBIC)
        cx, cy = (x1 + x2) / 2, (y1 + y2) / 2
        paste_x, paste_y = round(cx - tile.width / 2), round(cy - tile.height / 2)
    else:
        paste_x, paste_y = x1, y1

    canvas.paste(tile, (paste_x, paste_y), tile)


def _connector_endpoints(
    rect: tuple[int, int, int, int], flip_h: bool, flip_v: bool
) -> tuple[tuple[float, float], tuple[float, float]]:
    """Determines, via `flipH`/`flipV`, which bbox diagonal ("\\" or "/") a
    straight connector (`prstGeom prst="line"`/`"straightConnector1"`)
    joins. The default (no flip) is top-left -> bottom-right ("\\"), and if
    exactly one of `flipH`/`flipV` is on, the diagonal flips to the other
    ("/") — a diagonal is symmetric across the axes, so either `flipH` alone
    or `flipV` alone gives the same result, and turning both on (equivalent
    to a 180-degree rotation) goes back to the original (only the direction
    reverses; the two corners it passes through are the same) — hence the
    XOR test."""
    x1, y1, x2, y2 = rect
    if flip_h != flip_v:
        return (x2, y1), (x1, y2)  # "/"
    return (x1, y1), (x2, y2)  # "\"


def _draw_arrowhead(
    draw: ImageDraw.ImageDraw, tip: tuple[float, float], other_end: tuple[float, float], size: float, rgb: tuple[int, int, int]
) -> None:
    """Draws a filled triangular arrowhead pointing from `other_end` to
    `tip`. Doesn't distinguish the exact preset (triangle/stealth/diamond/
    oval etc.) and approximates all of them as a triangle (see
    `_has_arrowhead`)."""
    dx, dy = tip[0] - other_end[0], tip[1] - other_end[1]
    length = math.hypot(dx, dy)
    if length < 1e-6:
        return
    dx, dy = dx / length, dy / length
    px, py = -dy, dx  # unit vector perpendicular to the direction of travel
    base_x, base_y = tip[0] - dx * size, tip[1] - dy * size
    left = (base_x + px * size * 0.5, base_y + py * size * 0.5)
    right = (base_x - px * size * 0.5, base_y - py * size * 0.5)
    draw.polygon([tip, left, right], fill=rgb)


def _cell_fill_rgb(cell, theme_colors: list[str] | None) -> tuple[int, int, int] | None:
    fill = cell.fill
    if fill is None or fill.patternType != "solid":
        return None
    return _resolve_color(fill.fgColor, theme_colors)


def _border_segment(side, theme_colors: list[str] | None) -> tuple[int, tuple[int, int, int]] | None:
    """Returns one border side's (line width px, RGB). None if there's no
    style (the caller substitutes the default light-gray gridline)."""
    if side is None or side.style is None:
        return None
    width = _BORDER_WIDTH_PX.get(side.style, 1)
    rgb = _resolve_color(side.color, theme_colors) or (0, 0, 0)
    return width, rgb


def _wrap_lines(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    """Wraps by whitespace, but also splits by character when a single
    chunk with no spaces already exceeds `max_width` (e.g. a long URL/code
    string with no spaces)."""
    if max_width <= 0:
        return [text]
    words = text.split(" ")
    lines: list[str] = []
    cur = ""
    for word in words:
        trial = f"{cur} {word}".strip()
        if not trial:
            continue
        if draw.textlength(trial, font=font) <= max_width or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = word
        # force-split by character when even a single word is already too long
        while draw.textlength(cur, font=font) > max_width and len(cur) > 1:
            lo, hi = 1, len(cur)
            while lo < hi:
                mid = (lo + hi + 1) // 2
                if draw.textlength(cur[:mid], font=font) <= max_width:
                    lo = mid
                else:
                    hi = mid - 1
            lines.append(cur[:lo])
            cur = cur[lo:]
    if cur:
        lines.append(cur)
    return lines or [""]


def capture_table_image(
    ws,
    min_row: int,
    min_col: int,
    max_row: int,
    max_col: int,
    out_path: Path,
    padding: int = 1,
    xlsx_path: Path | None = None,
) -> tuple[int, int, int]:
    """Composites a table range into a PNG and saves it to `out_path`.

    `padding`: how many cells of margin to add outside the table's boundary
    when sizing the canvas (default 1). When an anchored image inside the
    table slightly overflows the table boundary (from cell-size estimation
    error, or a slight overflow present from the original), cropping the
    canvas exactly to the table size would clip that overflow — the margin
    is left blank with no gridlines/background/text, just providing room for
    an image to overflow into.

    `xlsx_path`: the original xlsx file's path. If given, shapes a person
    drew directly (highlight rectangles/ellipses, rotated shapes, arrows and
    other straight connectors, `<xdr:sp>`/`<xdr:cxnSp>`) are drawn too —
    since openpyxl never parses these shapes (see `_load_sheet_shapes`), the
    zip has to be reopened, which needs the path. `None` (the default)
    keeps the old behavior of not drawing shapes (backward compatible).
    Bent/curved connectors (`bentConnector*`/`curvedConnector*`) still
    aren't supported, since they need an adjustment-value-based path
    formula.

    Returns: (canvas width, canvas height, number of embedded images composited).
    """
    sheet_max_row = ws.max_row or max_row
    sheet_max_col = ws.max_column or max_col
    pad_min_row = max(1, min_row - padding)
    pad_max_row = min(sheet_max_row, max_row + padding)
    pad_min_col = max(1, min_col - padding)
    pad_max_col = min(sheet_max_col, max_col + padding)

    mdw_px = _sheet_mdw_px(ws)
    row_h = {r: _row_height_px(ws, r) for r in range(pad_min_row, pad_max_row + 1)}
    col_w = {c: _col_width_px(ws, c, mdw_px) for c in range(pad_min_col, pad_max_col + 1)}

    # row_off/col_off are keyed by absolute row/column index, padding
    # included — since the loops below that draw table content only iterate
    # over the original min/max range, the padding area is automatically
    # left as blank margin (white background), and it's actually used only
    # when an image overflows the table boundary.
    row_off = _cumulative_offsets(row_h, pad_min_row, pad_max_row)
    col_off = _cumulative_offsets(col_w, pad_min_col, pad_max_col)
    canvas_w, canvas_h = col_off[pad_max_col + 1], row_off[pad_max_row + 1]
    if canvas_w <= 0 or canvas_h <= 0 or canvas_w * canvas_h > MAX_CANVAS_PIXELS:
        return canvas_w, canvas_h, 0

    theme_colors = _parse_theme_colors(ws.parent)

    canvas = PILImage.new("RGB", (canvas_w, canvas_h), "white")
    draw = ImageDraw.Draw(canvas)

    # Filter down to merged cells overlapping the table range, and
    # precompute each merge anchor (top-left) -> its full pixel rectangle,
    # plus each cell coordinate -> the anchor of the merge it belongs to.
    # Internal merge boundary lines aren't drawn (so it looks like a single
    # real cell), and text/background are also drawn relative to the whole
    # merged rectangle.
    merge_anchor_of: dict[tuple[int, int], tuple[int, int]] = {}
    merge_rect_px: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    for m in ws.merged_cells.ranges:
        if m.max_row < min_row or m.min_row > max_row or m.max_col < min_col or m.min_col > max_col:
            continue
        anchor = (m.min_row, m.min_col)
        r1, c1 = max(m.min_row, min_row), max(m.min_col, min_col)
        r2, c2 = min(m.max_row, max_row), min(m.max_col, max_col)
        rect = (col_off[c1], row_off[r1], col_off.get(c2 + 1, canvas_w), row_off.get(r2 + 1, canvas_h))
        merge_rect_px[anchor] = rect
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                merge_anchor_of[(r, c)] = anchor

    # 1. Fill background color (must be drawn before text/gridlines so it doesn't cover them)
    painted_bg: set[tuple[int, int]] = set()
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            anchor = merge_anchor_of.get((r, c), (r, c))
            if anchor in painted_bg:
                continue
            painted_bg.add(anchor)
            rgb = _cell_fill_rgb(ws.cell(row=anchor[0], column=anchor[1]), theme_colors)
            if rgb is None:
                continue
            if anchor in merge_rect_px:
                rect = merge_rect_px[anchor]
            else:
                rect = (col_off[c], row_off[r], col_off[c + 1], row_off[r + 1])
            draw.rectangle(rect, fill=rgb)

    # 2. Gridlines — internal merge boundaries are skipped. Uses the actual
    # border formatting's width/color where present, otherwise the default
    # light-gray hairline.
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(row=r, column=c)
            anchor = merge_anchor_of.get((r, c))
            x1, y1, x2, y2 = col_off[c], row_off[r], col_off[c + 1], row_off[r + 1]

            top_is_internal = anchor is not None and merge_anchor_of.get((r - 1, c)) == anchor
            if not top_is_internal:
                seg = _border_segment(cell.border.top, theme_colors)
                width, rgb = seg if seg else (1, _GRID_RGB)
                draw.line([(x1, y1), (x2, y1)], fill=rgb, width=width)

            left_is_internal = anchor is not None and merge_anchor_of.get((r, c - 1)) == anchor
            if not left_is_internal:
                seg = _border_segment(cell.border.left, theme_colors)
                width, rgb = seg if seg else (1, _GRID_RGB)
                draw.line([(x1, y1), (x1, y2)], fill=rgb, width=width)

            if r == max_row:
                seg = _border_segment(cell.border.bottom, theme_colors)
                width, rgb = seg if seg else (1, _GRID_RGB)
                draw.line([(x1, y2), (x2, y2)], fill=rgb, width=width)
            if c == max_col:
                seg = _border_segment(cell.border.right, theme_colors)
                width, rgb = seg if seg else (1, _GRID_RGB)
                draw.line([(x2, y1), (x2, y2)], fill=rgb, width=width)

    # 3. Text — a merged cell aligns/wraps relative to the whole rectangle.
    drawn_text: set[tuple[int, int]] = set()
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            anchor = merge_anchor_of.get((r, c), (r, c))
            if anchor in drawn_text:
                continue
            drawn_text.add(anchor)
            cell = ws.cell(row=anchor[0], column=anchor[1])
            value = cell.value
            if value is None:
                continue
            # Using str(value) alone ignores number_format, so a '0.00'
            # format cell prints its double-precision float value as-is (15
            # decimal digits) or a '0%' percent format cell shows a number
            # 100x wrong (see number_format.py).
            text = format_cell_display(value, cell.number_format)[:MAX_CELL_TEXT_CHARS]

            if anchor in merge_rect_px:
                x1, y1, x2, y2 = merge_rect_px[anchor]
            else:
                x1, y1, x2, y2 = col_off[c], row_off[r], col_off[c + 1], row_off[r + 1]
            box_w, box_h = max(1, x2 - x1 - 4), max(1, y2 - y1 - 2)

            font_pt = cell.font.sz or 11
            font_px = min(40, max(8, round(font_pt * DPI / 72 * 0.85)))  # 0.85: correction since using pt as-is looks somewhat too large relative to the cell
            font = _load_capture_font(font_px, bold=bool(cell.font.b))
            color = _resolve_color(cell.font.color, theme_colors) or _DEFAULT_TEXT_RGB

            wrap = bool(cell.alignment.wrap_text)
            raw_lines = text.split("\n")
            lines: list[str] = []
            for raw_line in raw_lines:
                lines.extend(_wrap_lines(draw, raw_line, font, box_w) if wrap else [raw_line])

            line_h = round(font_px * 1.3)
            block_h = line_h * len(lines)
            valign = cell.alignment.vertical or "bottom"
            if valign == "center":
                y = y1 + 1 + max(0, (box_h - block_h) // 2)
            elif valign == "top":
                y = y1 + 1
            else:  # bottom/default — Excel's default vertical alignment is bottom
                y = y1 + 1 + max(0, box_h - block_h)

            halign = cell.alignment.horizontal or ("right" if isinstance(value, (int, float)) else "left")
            for line in lines:
                line_w = draw.textlength(line, font=font)
                if halign == "center":
                    x = x1 + 2 + max(0, (box_w - line_w) / 2)
                elif halign == "right":
                    x = x1 + 2 + max(0, box_w - line_w)
                else:  # left/general/other
                    x = x1 + 2
                draw.text((x, y), line, fill=color, font=font)
                y += line_h

    # 4. Actual embedded images inside the table — drawn over text/grid
    # (the photo has to take priority for the table content to be legible).
    pasted = 0
    for embedded in getattr(ws, "_images", []):
        anchor = embedded.anchor
        fr = getattr(anchor, "_from", None)
        if fr is None:
            continue
        r, c = fr.row + 1, fr.col + 1
        if not (min_row <= r <= max_row and min_col <= c <= max_col):
            continue
        try:
            pic = PILImage.open(io.BytesIO(embedded._data())).convert("RGBA")
        except Exception:  # noqa: BLE001 — a format PIL can't open (an exceptional case) is skipped
            continue

        x = col_off[c] + _emu_to_px(fr.colOff)
        y = row_off[r] + _emu_to_px(fr.rowOff)

        to = getattr(anchor, "to", None)
        # For a TwoCellAnchor with editAs="oneCell" (or "absolute"), even if
        # `to` is present, its value isn't the basis for size — per the
        # OOXML spec this means "move, but don't resize to fit the cell," so
        # `to` is just a coordinate Excel fills in as a rough estimate for
        # screen refresh, not the real render size. In this case the real
        # size comes from the shape's own transform extent
        # (`anchor.pic.spPr.xfrm.ext`) — exactly the same concept as
        # OneCellAnchor's `ext`.
        edit_as = getattr(anchor, "editAs", None)
        fixed_size_anchor = to is None or edit_as in ("oneCell", "absolute")

        if to is not None and not fixed_size_anchor:
            # TwoCellAnchor, editAs default ("twoCell") — the case that's
            # genuinely resized to fit the cell. Computing both corners
            # directly from the start/end cell+offset keeps row-height
            # estimation error from accumulating.
            to_row, to_col = to.row + 1, to.col + 1
            # If it points slightly outside the table range to a key not in
            # row_off/col_off, clamp to the canvas edge (avoid a dict
            # lookup failure).
            x2 = col_off[to_col] + _emu_to_px(to.colOff) if to_col in col_off else canvas_w
            y2 = row_off[to_row] + _emu_to_px(to.rowOff) if to_row in row_off else canvas_h
            w, h = max(1, x2 - x), max(1, y2 - y)
        else:
            # OneCellAnchor, or a TwoCellAnchor with editAs="oneCell"/"absolute"
            # — size comes from the shape's own transform extent (independent of
            # cell size, fixed size).
            ext = getattr(anchor, "ext", None)  # present directly for OneCellAnchor
            if ext is None:
                try:
                    ext = anchor.pic.spPr.xfrm.ext  # TwoCellAnchor(oneCell/absolute)
                except AttributeError:
                    ext = None
            if ext is not None:
                w, h = max(1, _emu_to_px(ext.cx)), max(1, _emu_to_px(ext.cy))
            else:
                w = round(embedded.width) if embedded.width else pic.width
                h = round(embedded.height) if embedded.height else pic.height

        pic = pic.resize((w, h))
        canvas.paste(pic, (x, y), pic if pic.mode == "RGBA" else None)
        pasted += 1

    # 5. Shapes a person drew directly (highlight rectangle/ellipse boxes,
    # arrows, etc.) — drawn after (on top of) embedded images, since they're
    # often used as annotations on a photo. A shape whose fill and outline
    # both couldn't be read (only a theme style reference with no override,
    # see `_shape_style`) is silently skipped rather than inventing a color.
    # Bent/curved connectors (bentConnector/curvedConnector) still need an
    # adjustment-value-based path formula and are out of scope — only
    # straight connectors are supported.
    if xlsx_path is not None:
        for anchor, rect, elem, kind in _flatten_shape_entries(xlsx_path, ws, col_off, row_off, canvas_w, canvas_h):
            frm = anchor.find(f"{_XDR_NS}from")
            if frm is None:
                continue
            fr = int(frm.findtext(f"{_XDR_NS}row", "0")) + 1
            fc = int(frm.findtext(f"{_XDR_NS}col", "0")) + 1
            if not (min_row <= fr <= max_row and min_col <= fc <= max_col):
                continue  # not a shape anchored within this table range

            prst, fill, outline, rotation_deg = _shape_style(elem, theme_colors)

            if kind == "sp":
                text_lines = _shape_text_lines(elem, theme_colors)
                if fill is None and outline is None and not text_lines:
                    continue
                _paint_shape(canvas, rect, prst, fill, outline, rotation_deg, text_lines)
                continue

            # kind == "cxnSp": only straight connectors are supported (see the comment above for bent/curved).
            if prst.startswith(_BENT_OR_CURVED_CONNECTOR_PREFIXES) or outline is None:
                continue
            sppr = elem.find(f"{_XDR_NS}spPr")
            flip_h, flip_v, _ = _xfrm_transform(sppr)
            p1, p2 = _connector_endpoints(rect, flip_h, flip_v)
            (r, g, b), width = outline
            draw.line([p1, p2], fill=(r, g, b), width=width)
            ln = sppr.find(f"{_THEME_NS}ln") if sppr is not None else None
            arrow_size = max(6.0, width * 3.0)
            if _has_arrowhead(ln, "tailEnd"):
                _draw_arrowhead(draw, p1, p2, arrow_size, (r, g, b))
            if _has_arrowhead(ln, "headEnd"):
                _draw_arrowhead(draw, p2, p1, arrow_size, (r, g, b))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(out_path)
    return canvas_w, canvas_h, pasted


def image_anchor_rect(
    img,
    row_off: dict[int, int],
    col_off: dict[int, int],
    canvas_w: int,
    canvas_h: int,
) -> tuple[int, int, int, int] | None:
    """Resolves an openpyxl `Image`'s (an entry in `ws._images`) anchor to an
    absolute pixel rectangle.

    Pulled out into one function because `capture_table_image` (embedded
    images inside a table) and `annotate_standalone_image` (the image
    itself, and other overlapping standalone images) both repeat the same
    TwoCellAnchor (genuinely resized to the cell) / OneCellAnchor·
    editAs="oneCell"·"absolute" (fixed size, `to` is just an approximate
    coordinate for screen refresh) distinction — see either function's
    original docstring for the detailed rules. Exposed across the module
    boundary because `extractors/xlsx.py` needs the same rule to resolve
    these anchors when checking whether standalone images outside any table
    overlap each other (e.g. a small icon photo placed separately on top of
    another photo). None if the anchor's start cell is outside the
    `row_off`/`col_off` range."""
    anchor = img.anchor
    fr = getattr(anchor, "_from", None)
    if fr is None:
        return None
    r0, c0 = fr.row + 1, fr.col + 1
    if r0 not in row_off or c0 not in col_off:
        return None
    x = col_off[c0] + _emu_to_px(fr.colOff)
    y = row_off[r0] + _emu_to_px(fr.rowOff)

    to = getattr(anchor, "to", None)
    edit_as = getattr(anchor, "editAs", None)
    # With editAs="oneCell" (or "absolute"), `to` is just an approximate
    # coordinate for screen refresh, not the basis for size — the same rule
    # as capture_table_image's embedded-image handling.
    fixed_size_anchor = to is None or edit_as in ("oneCell", "absolute")
    if to is not None and not fixed_size_anchor:
        to_row, to_col = to.row + 1, to.col + 1
        x2 = col_off[to_col] + _emu_to_px(to.colOff) if to_col in col_off else canvas_w
        y2 = row_off[to_row] + _emu_to_px(to.rowOff) if to_row in row_off else canvas_h
        w, h = max(1, x2 - x), max(1, y2 - y)
    else:
        ext = getattr(anchor, "ext", None)
        if ext is None:
            try:
                ext = anchor.pic.spPr.xfrm.ext
            except AttributeError:
                ext = None
        if ext is not None:
            w, h = max(1, _emu_to_px(ext.cx)), max(1, _emu_to_px(ext.cy))
        else:
            w = round(img.width) if img.width else 100
            h = round(img.height) if img.height else 100
    return x, y, x + w, y + h


def annotate_standalone_image(
    ws,
    img,
    image_bytes: bytes,
    xlsx_path: Path,
    row_off: dict[int, int],
    col_off: dict[int, int],
    canvas_w: int,
    canvas_h: int,
    out_path: Path,
    overlay_images: list[tuple[bytes, tuple[int, int, int, int]]] | None = None,
) -> tuple[int, int, int, int] | None:
    """Composites annotations a person placed on top of a standalone Image
    that isn't inside any Table range (highlight rectangles/ellipses,
    arrows and other straight connectors, text labels), along with any
    other standalone image placed separately on top of it (e.g. a driver
    icon photo), and saves the result as a PNG.

    `capture_table_image` already composites shapes and embedded images
    inside a table range (the "Option A: absorb" design, see the module
    docstring), but until now a standalone image outside any table only had
    its raw bytes saved by `extractors/xlsx.py`, so any annotation placed on
    top of it (a red highlight box, a yellow arrow/label, etc.) or another
    photo on top (e.g. a driver icon, confirmed in a screw fastening-
    strength review document) was entirely lost — this function fills that
    gap.

    `image_bytes`: takes the raw bytes the caller already read, as-is —
    openpyxl's `Image._data()` consumes and closes its internal file handle,
    so **calling it a second time dies with `ValueError: I/O operation on
    closed file`** (confirmed). The caller has already saved the original
    once via `img._data()` before calling this function, so `img._data()`
    must not be called again here.

    `overlay_images`: a list of (raw bytes, absolute pixel rectangle) for
    **other standalone images** that overlap this one — `extractors/xlsx.py`
    pre-checks whether standalone images outside any table overlap each
    other using `image_anchor_rect` and passes the result in (unlike
    `xdr:sp`/`xdr:cxnSp` shapes, which openpyxl can't parse, a photo placed
    separately on top of another photo is already fully captured by
    `ws._images`, so the caller only needs to compute anchor overlap instead
    of re-scanning raw XML like for shapes). Images in this list don't
    become a separate Image node (the same "Option A: absorb" principle —
    here another standalone image is the absorption target instead of a
    table) — this function composites and preserves them instead.

    An image anchor and a shape anchor can reference different cell ranges
    (e.g. a highlight box starting a cell or two before the photo), so
    there's no basis to scope things to "inside that table range" the way
    table capture does — instead, every anchor is resolved into the same
    absolute pixel coordinates via `row_off`/`col_off` (the sheet-wide
    coordinate system, computed once per sheet by `build_sheet_grid` and
    reused per image), and only the shapes/images that actually overlap the
    image's own bbox are selected and composited.

    Returns `None` if there are no overlapping shapes and no overlapping
    images at all — this lets the caller keep using the raw bytes it already
    saved, so the overwhelming majority of images with no annotation stay as
    the original with no unnecessary re-encoding. The return value is
    (canvas width, canvas height, number of shapes composited, number of
    images composited) — these two counts are kept separate because
    `extractors/xlsx.py` needs to record the shape-annotation count
    (`annotation_shape_count`, existing schema) and the absorbed-image count
    (`absorbed_image_count`, new) as separate properties."""
    image_rect = image_anchor_rect(img, row_off, col_off, canvas_w, canvas_h)
    if image_rect is None:
        return None

    try:
        pic = PILImage.open(io.BytesIO(image_bytes)).convert("RGBA")
    except Exception:  # noqa: BLE001 — e.g. a corrupted embedded image
        return None
    w, h = image_rect[2] - image_rect[0], image_rect[3] - image_rect[1]

    theme_colors = _parse_theme_colors(ws.parent)
    overlaps: list[tuple[ET.Element, str, tuple[int, int, int, int]]] = []
    for _sh_anchor, rect, elem, kind in _flatten_shape_entries(xlsx_path, ws, col_off, row_off, canvas_w, canvas_h):
        if not rects_overlap(rect, image_rect):
            continue
        overlaps.append((elem, kind, rect))

    # Other overlapping standalone images — the caller already did the
    # overlap check before passing these in, but this function verifies it
    # again itself too (so that a caller mistake passing in a non-
    # overlapping one is silently ignored).
    valid_overlay_images = [
        (data, rect) for data, rect in (overlay_images or []) if rects_overlap(rect, image_rect)
    ]

    if not overlaps and not valid_overlay_images:
        return None

    all_rects = [image_rect] + [r for _, _, r in overlaps] + [r for _, r in valid_overlay_images]
    ux1 = min(r[0] for r in all_rects)
    uy1 = min(r[1] for r in all_rects)
    ux2 = max(r[2] for r in all_rects)
    uy2 = max(r[3] for r in all_rects)
    out_w, out_h = ux2 - ux1, uy2 - uy1
    if out_w <= 0 or out_h <= 0 or out_w * out_h > MAX_CANVAS_PIXELS:
        return None

    canvas = PILImage.new("RGB", (out_w, out_h), "white")
    pic_resized = pic.resize((w, h))
    canvas.paste(pic_resized, (image_rect[0] - ux1, image_rect[1] - uy1), pic_resized)

    # Other overlapping standalone images — pasted before shapes
    # (rectangles/arrows etc.). Same ordering as `capture_table_image`
    # drawing embedded images (step 4) before shapes (step 5) — a
    # person-drawn highlight shape is often the topmost layer over the
    # photos.
    image_drawn = 0
    for data, rect in valid_overlay_images:
        try:
            overlay_pic = PILImage.open(io.BytesIO(data)).convert("RGBA")
        except Exception:  # noqa: BLE001 — e.g. a corrupted embedded image, skip just that one
            continue
        ow, oh = max(1, rect[2] - rect[0]), max(1, rect[3] - rect[1])
        overlay_pic = overlay_pic.resize((ow, oh))
        canvas.paste(overlay_pic, (rect[0] - ux1, rect[1] - uy1), overlay_pic)
        image_drawn += 1

    draw = ImageDraw.Draw(canvas)

    shape_drawn = 0
    for elem, kind, rect in overlaps:
        local_rect = (rect[0] - ux1, rect[1] - uy1, rect[2] - ux1, rect[3] - uy1)
        prst, fill, outline, rotation_deg = _shape_style(elem, theme_colors)

        if kind == "sp":
            text_lines = _shape_text_lines(elem, theme_colors)
            if fill is None and outline is None and not text_lines:
                continue  # a shape whose color and text couldn't both be read — no guessing
            _paint_shape(canvas, local_rect, prst, fill, outline, rotation_deg, text_lines)
            shape_drawn += 1
            continue

        # kind == "cxnSp": only straight connectors are supported (same scope
        # constraint as capture_table_image, see that comment).
        if prst.startswith(_BENT_OR_CURVED_CONNECTOR_PREFIXES) or outline is None:
            continue
        sppr = elem.find(f"{_XDR_NS}spPr")
        flip_h, flip_v, _ = _xfrm_transform(sppr)
        p1, p2 = _connector_endpoints(local_rect, flip_h, flip_v)
        (r, g, b), width = outline
        draw.line([p1, p2], fill=(r, g, b), width=width)
        ln = sppr.find(f"{_THEME_NS}ln") if sppr is not None else None
        arrow_size = max(6.0, width * 3.0)
        if _has_arrowhead(ln, "tailEnd"):
            _draw_arrowhead(draw, p1, p2, arrow_size, (r, g, b))
        if _has_arrowhead(ln, "headEnd"):
            _draw_arrowhead(draw, p2, p1, arrow_size, (r, g, b))
        shape_drawn += 1

    if shape_drawn == 0 and image_drawn == 0:
        return None

    out_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(out_path)
    return out_w, out_h, shape_drawn, image_drawn
